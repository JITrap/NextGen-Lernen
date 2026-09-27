"""
Excel-Planrechnung (mit Formeln) für den Businessplan – erzeugt aus annahmen.json + daten.json.
Blätter: Deckblatt, Annahmen, Investition, Finanzierung, Mitglieder, Umsatz, Kosten, Liquiditaet, GuV, Kennzahlen, Szenarien.
Konventionen: blaue Schrift = Eingabe, schwarz = Formel, grün = Verweis auf anderes Blatt, gelbe Füllung = vor dem Bankgespräch ausfüllen/prüfen.
"""
from __future__ import annotations

import datetime as dt
import json
from typing import Any

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.worksheet import Worksheet

import model as M

FONT = 'Arial'
BLUE = Font(name=FONT, size=10, color='0000FF')
BLACK = Font(name=FONT, size=10)
GREEN = Font(name=FONT, size=10, color='008000')
BOLD = Font(name=FONT, size=10, bold=True)
H1 = Font(name=FONT, size=14, bold=True)
H2 = Font(name=FONT, size=11, bold=True)
GREY = Font(name=FONT, size=9, italic=True, color='666666')
YELLOW = PatternFill('solid', fgColor='FFFF00')
HEAD = PatternFill('solid', fgColor='D9E1F2')
SUB = PatternFill('solid', fgColor='F2F2F2')
THIN = Side(style='thin', color='999999')
TOP = Border(top=THIN)
EUR = '#,##0;[Red]-#,##0;-'
EUR2 = '#,##0.00;[Red]-#,##0.00;-'
PCT = '0.0%'
PCT2 = '0.00%'
NUM = '#,##0;[Red]-#,##0;-'
NUM1 = '#,##0.0'

MONATE = M.MONATE


class Names:
    """Sammelt definierte Namen (Name → 'Blatt'!$B$5)."""

    def __init__(self, wb: Workbook):
        self.wb = wb

    def add(self, name: str, ws: Worksheet, cell: str) -> None:
        self.wb.defined_names[name] = DefinedName(name, attr_text=f"'{ws.title}'!{absref(cell)}")

    def add_range(self, name: str, ws: Worksheet, rng: str) -> None:
        a, b = rng.split(':')
        self.wb.defined_names[name] = DefinedName(name, attr_text=f"'{ws.title}'!{absref(a)}:{absref(b)}")


def absref(cell: str) -> str:
    col = ''.join(ch for ch in cell if ch.isalpha())
    row = ''.join(ch for ch in cell if ch.isdigit())
    return f'${col}${row}'


def setw(ws: Worksheet, widths: dict[str, float]) -> None:
    for c, w in widths.items():
        ws.column_dimensions[c].width = w


def put(ws: Worksheet, cell: str, value: Any, font: Font = BLACK, fmt: str | None = None, fill: PatternFill | None = None, comment: str | None = None, align: str | None = None, bold: bool = False) -> None:
    c = ws[cell]
    c.value = value
    c.font = Font(name=FONT, size=font.size, bold=bold or font.bold, italic=font.italic, color=font.color)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if comment:
        c.comment = Comment(comment, 'Planung')
    if align:
        c.alignment = Alignment(horizontal=align, wrap_text=True)


# ---------------------------------------------------------------------------
def write_excel(a: dict, daten: dict, path: str, szenario_werte: dict[str, dict] | None = None, stand: str = '') -> None:
    wb = Workbook()
    names = Names(wb)
    P = a['projekt']
    V = P['vorlauf_monate']
    mi = a['miete']
    t = a['tarife']
    nu = a['nebenumsatz']
    mg = a['mitglieder']
    pe = a['personal']
    sk = a['sachkosten_monat']
    vo = a['vorlauf']
    inv_a = a['investition']
    fin = a['finanzierung']
    st = a['steuern']
    monate_all = list(range(-V, 0)) + list(range(1, MONATE + 1))
    stand = stand or dt.date.today().strftime('%d.%m.%Y')

    # ======================= Deckblatt =======================
    ws = wb.active
    ws.title = 'Deckblatt'
    setw(ws, {'A': 3, 'B': 34, 'C': 70})
    put(ws, 'B2', f"Finanzplanung {P['name']}", H1)
    put(ws, 'B3', f"Businessplan-Anlage: Kapitalbedarf, Finanzierungsplan, Mitglieder-/Umsatz-/Kostenplanung, Liquiditätsplan, Rentabilitätsvorschau (5 Jahre)", BLACK)
    put(ws, 'B4', f"Stand: {stand} · Grundlage: Planung „{daten.get('projekt', {}).get('name', 'No.1 (überarbeitet)')}“ aus GymPlanner ({P['flaeche_brutto_m2']:.0f} m² Bruttofläche), Kaltmiete {mi['kalt_eur_m2']:.2f} €/m², Anlaufreserve für 24 Monate", GREY)
    put(ws, 'B6', 'Legende', H2)
    put(ws, 'B7', 'Blaue Schrift', BLUE); put(ws, 'C7', 'Eingabewert / Annahme – darf geändert werden (Blatt „Annahmen“, „Investition“)')
    put(ws, 'B8', 'Schwarze Schrift', BLACK); put(ws, 'C8', 'Formel – nicht überschreiben')
    put(ws, 'B9', 'Grüne Schrift', GREEN); put(ws, 'C9', 'Verweis auf ein anderes Blatt')
    put(ws, 'B10', 'Gelbe Füllung', BLACK, fill=YELLOW); put(ws, 'C10', 'Vor dem Bankgespräch prüfen bzw. persönlich ausfüllen (Eigenkapital, Gründerdaten, Eröffnungstermin)')
    put(ws, 'B12', 'Blätter', H2)
    rows = [
        ('Annahmen', 'Alle Eingaben: Fläche, Miete, Tarife, Mitgliederentwicklung, Personal, Sachkosten, Finanzierung, Steuern, Szenario-Schalter'),
        ('Investition', 'Kapitalbedarf: Geräte je Bereich (Stückliste GymPlanner), Ausbau, Sanitär, Planung, Unvorhergesehenes, Kaution, Gründung, Anlaufreserve; Zahlungsplan Vorlaufmonate'),
        ('Finanzierung', 'Mittelherkunft (Eigenkapital, Darlehen), Tilgungspläne je Darlehen (monatlich), Kapitaldienst je Jahr'),
        ('Mitglieder', 'Neuzugänge, Kündigungen, Bestand je Monat (60 Monate), Spitzenlast gegen Kapazität'),
        ('Umsatz', 'Beiträge, Servicepauschale, Aufnahmegebühren, Nebenumsätze – netto je Monat, Umsatzsteuer'),
        ('Kosten', 'Personal (Stellenplan), Raumkosten, Marketing, sonstige Sachkosten je Monat inkl. Vorlaufphase, Vorsteuer'),
        ('Liquiditaet', 'Monatlicher Liquiditätsplan von der Ausbauphase bis Monat 60 inkl. Umsatzsteuer-Zahllast/-Erstattung, Kapitaldienst, Steuern'),
        ('GuV', 'Rentabilitätsvorschau Jahr 1–5, Kapitaldienstfähigkeit (DSCR), Abschreibungen, Steuern mit Verlustvortrag'),
        ('Kennzahlen', 'Break-even (Mitglieder/Monat), Mindestliquidität, Reservebedarf, Quoten'),
        ('Szenarien', 'Vergleich pessimistisch / Basis / optimistisch (Werte, erzeugt durch Umschalten der Zelle „Szenario“)'),
    ]
    for i, (n, d) in enumerate(rows):
        put(ws, f'B{13 + i}', n, BOLD); put(ws, f'C{13 + i}', d, BLACK, align='left')
    put(ws, 'B25', 'Hinweise', H2)
    hints = [
        'Alle Beträge in Euro netto (ohne Umsatzsteuer), sofern nicht „brutto“ angegeben. Mitgliedsbeiträge sind Bruttopreise inkl. 19 % USt; die Planung rechnet sie auf netto um.',
        'Monatsraster: Monate −5 … −1 = Vorlauf (Ausbau, Vorverkauf), Monat 1 = Eröffnungsmonat, Jahr 1 = Vorlauf + Monate 1–12 (Rumpf-/erstes Geschäftsjahr).',
        'Investitionsbeträge stammen aus der Kostenkalkulation der Planungssoftware (Stückliste mit recherchierten Netto-Listenpreisen, Ausbau-Richtwerte je m²). Sie sind vor Vertragsabschluss durch Angebote zu ersetzen.',
        'Die Anlaufreserve (Betriebsmittel) deckt die kumulierten Anlaufverluste der ersten 24 Monate inkl. Sicherheitspuffer; Blatt „Kennzahlen“ zeigt den berechneten Bedarf gegenüber dem angesetzten Betrag.',
        'Darlehenskonditionen sind Planannahmen auf Basis der veröffentlichten Programmbedingungen (KfW/L-Bank/Bürgschaftsbank); die Hausbank legt die Preisklasse fest.',
    ]
    for i, h in enumerate(hints):
        put(ws, f'B{26 + i}', f'{i + 1}.', BLACK); put(ws, f'C{26 + i}', h, BLACK, align='left')
        ws.row_dimensions[26 + i].height = 30

    # ======================= Annahmen =======================
    an = wb.create_sheet('Annahmen')
    setw(an, {'A': 3, 'B': 52, 'C': 16, 'D': 12, 'E': 60})
    put(an, 'B1', 'Annahmen (Eingaben)', H1)
    put(an, 'B2', 'Blau = Eingabe. Gelb = vor dem Bankgespräch ausfüllen/prüfen. Quellen und Herleitung siehe Businessplan Kapitel „Finanzplanung“ und Spalte E.', GREY)
    r = 4

    def head(title: str) -> None:
        nonlocal r
        put(an, f'B{r}', title, H2, fill=HEAD); an[f'C{r}'].fill = HEAD; an[f'D{r}'].fill = HEAD; an[f'E{r}'].fill = HEAD
        r += 1

    def inp(name: str, label: str, value: Any, fmt: str, unit: str = '', note: str = '', yellow: bool = False, formula: bool = False) -> str:
        nonlocal r
        put(an, f'B{r}', label)
        put(an, f'C{r}', value, BLACK if formula else BLUE, fmt, YELLOW if yellow else None)
        put(an, f'D{r}', unit, GREY)
        put(an, f'E{r}', note, GREY, align='left')
        names.add(name, an, f'C{r}')
        cell = f'C{r}'
        r += 1
        return cell

    head('Projekt')
    inp('Flaeche', 'Bruttofläche (Mietfläche)', P['flaeche_brutto_m2'], NUM1, 'm²', 'aus Planung No.1 (überarbeitet): Hallenaußenmaß')
    inp('Flaeche_Netto', 'Nettofläche', P['flaeche_netto_m2'], NUM1, 'm²', 'Innenfläche ohne Wände')
    inp('Trainingsflaeche', 'Trainingsfläche (Zonen + Kursraum)', P['trainingsflaeche_m2'], NUM1, 'm²', 'Basis der Personenzahl (9 m² je Person)')
    inp('Personen', 'Gleichzeitige Nutzer (Kapazität)', P['personen_kapazitaet'], NUM, 'Pers.', 'Trainingsfläche ÷ 9 m²')
    inp('Vorlauf', 'Vorlaufmonate (Ausbau + Vorverkauf)', V, NUM, 'Monate', 'Monate −5 … −1 vor Eröffnung')
    c_eroeff = inp('Eroeffnung', 'Eröffnungstermin (geplant)', dt.datetime.strptime(P['eroeffnung'], '%Y-%m-%d'), 'DD.MM.YYYY', '', 'Monat 1 der Planung', yellow=True)
    inp('Szenario', 'Szenario (1 = pessimistisch, 2 = Basis, 3 = optimistisch)', 2, '0', '', 'schaltet Neuzugänge, Kündigungsquote und Beitragsniveau um (Tabelle unten)', yellow=True)

    head('Miete und Nebenkosten')
    inp('Miete_kalt', 'Kaltmiete', mi['kalt_eur_m2'], EUR2, '€/m²/Monat', 'Vorgabe: 8,50 €/m² (Hallen-/Sonderfläche, Region Bodensee/Stuttgart)')
    inp('BK', 'Betriebskosten (Grundsteuer, Versicherung, Müll, Hausmeister)', mi['betriebskosten_eur_m2'], EUR2, '€/m²/Monat', '')
    inp('Energie', 'Energie und Wasser (Strom, Wärme, Wasser/Abwasser)', mi['energie_wasser_eur_m2'], EUR2, '€/m²/Monat', 'inkl. Sauna/Dampfbad und Lüftung; erst ab Eröffnung')
    inp('Kaution_Monate', 'Kaution', mi['kaution_monate'], NUM, 'Monatsmieten', 'inkl. Nebenkosten (Investition)')
    inp('Mietfrei', 'Mietfreie Monate während des Ausbaus', mi['mietfreie_monate_ausbau'], NUM, 'Monate', 'Verhandlungsziel Mietvertrag; danach Miete + Betriebskosten ab Übergabe')
    inp('Index_pa', 'Mietindexierung je Jahr', mi['indexierung_prozent_pa'] / 100, PCT, '% p. a.', 'Indexmiete, ab Jahr 2')
    inp('UST_Miete', 'Vermieter optiert zur Umsatzsteuer (1 = ja, 0 = nein)', 1 if mi['ust_auf_miete'] else 0, '0', '', 'bei 1 ist die USt auf Miete/BK als Vorsteuer abziehbar')

    head('Tarife (brutto inkl. USt) und Mitgliedermix')
    inp('Beitrag_Basic', f"{t['basic']['name']}", t['basic']['beitrag_brutto'], EUR2, '€/Monat', '')
    inp('Anteil_Basic', 'Anteil Basic am Bestand', t['basic']['anteil'], PCT, '', '')
    inp('Beitrag_Premium', f"{t['premium']['name']}", t['premium']['beitrag_brutto'], EUR2, '€/Monat', '')
    inp('Anteil_Premium', 'Anteil Premium am Bestand', t['premium']['anteil'], PCT, '', '')
    inp('Beitrag_Student', f"{t['student']['name']}", t['student']['beitrag_brutto'], EUR2, '€/Monat', '')
    inp('Anteil_Student', 'Anteil Student/Azubi/Senior', t['student']['anteil'], PCT, '', 'Summe der Anteile = 100 %')
    inp('Service_Halbjahr', 'Servicepauschale je Halbjahr (brutto)', t.get('servicepauschale_halbjahr_brutto', 0), EUR2, '€', 'branchenüblich (Trainings-/Servicepauschale), 0 = keine')
    inp('Aufnahme', 'Aufnahmegebühr (brutto)', t['aufnahmegebuehr_brutto'], EUR2, '€', 'einmalig je Neuzugang')
    inp('Aufnahme_Anteil', 'Anteil Neuzugänge, die Aufnahmegebühr zahlen', t['aufnahmegebuehr_anteil_zahlend'], PCT, '', 'Rest über Aktionen erlassen')
    inp('Beitrag_Erhoehung', 'Beitragsanpassung ab Jahr 3', t.get('beitragserhoehung_prozent_pa', 0) / 100, PCT, '% p. a.', 'auf den Bestand (Indexklausel)')
    inp('UST', 'Umsatzsteuersatz', t['ust_satz'], PCT, '', 'Fitnessleistungen 19 %')
    c_arpu = inp('ARPU_brutto', 'Ø Beitrag je Mitglied und Monat (brutto, gewichteter Mix, Szenario)', '=(Beitrag_Basic*Anteil_Basic+Beitrag_Premium*Anteil_Premium+Beitrag_Student*Anteil_Student)*Beitrag_Faktor', EUR2, '€/Monat', 'Formel', formula=True)

    head('Nebenumsätze je Mitglied und Monat (brutto)')
    inp('Neben_Getraenke', 'Getränke, Shakes, Riegel', nu['getraenke_shakes_brutto_je_mitglied_monat'], EUR2, '€', '')
    inp('Neben_WE', 'Wareneinsatz Getränke/Shakes', nu['getraenke_wareneinsatz_prozent'], PCT, '% vom Netto', '')
    inp('Neben_Wellness', 'Solarium, Wellness-Zusatz, Getränke-Flat-Upgrade', nu['solarium_wellness_zusatz_brutto_je_mitglied_monat'], EUR2, '€', '')
    inp('Neben_PT', 'Personal Training, Zusatzkurse, Ernährungsberatung (Provisionsanteil)', nu['personal_training_kurse_brutto_je_mitglied_monat'], EUR2, '€', '')
    inp('Neben_Merch', 'Merchandise, Sonstiges', nu['merchandise_sonstiges_brutto_je_mitglied_monat'], EUR2, '€', '')

    head('Mitgliederentwicklung')
    inp('Vorverkauf', 'Mitglieder zur Eröffnung (Vorverkauf, 3 Monate)', mg['vorverkauf_bestand_eroeffnung'], NUM, 'Mitgl.', 'Basis-Szenario, wird mit Szenario-Faktor skaliert')
    ph = mg['neuzugaenge_je_monat']
    inp('Neu_1', f"Neuzugänge je Monat, Monate {ph[0]['von']}–{ph[0]['bis']}", ph[0]['n'], NUM, 'Mitgl./Monat', 'Eröffnungsphase')
    inp('Neu_2', f"Neuzugänge je Monat, Monate {ph[1]['von']}–{ph[1]['bis']}", ph[1]['n'], NUM, 'Mitgl./Monat', '')
    inp('Neu_3', f"Neuzugänge je Monat, Monate {ph[2]['von']}–{ph[2]['bis']}", ph[2]['n'], NUM, 'Mitgl./Monat', '')
    inp('Neu_4', f"Neuzugänge je Monat, Monate {ph[3]['von']}–{ph[3]['bis']}", ph[3]['n'], NUM, 'Mitgl./Monat', '')
    inp('Neu_5', f"Neuzugänge je Monat, Monate {ph[4]['von']}–{ph[4]['bis']}", ph[4]['n'], NUM, 'Mitgl./Monat', '')
    inp('Kapazitaet_Max', 'Maximaler Mitgliederbestand (Obergrenze)', mg['kapazitaet_max_mitglieder'], NUM, 'Mitgl.', f"≈ {mg['kapazitaet_max_mitglieder'] / P['flaeche_brutto_m2']:.2f} Mitglieder je m²")
    inp('Spitzenlast', 'Anteil gleichzeitig anwesender Mitglieder zur Spitzenzeit', mg['spitzenlast_anteil_gleichzeitig'], PCT, '', 'Prüfung gegen Personen-Kapazität (Blatt Mitglieder)')

    head('Szenarien (Faktoren)')
    put(an, f'B{r}', 'Szenario', BOLD); put(an, f'C{r}', 'Neuzugänge ×', BOLD); put(an, f'D{r}', 'Kündigung/Monat', BOLD); put(an, f'E{r}', 'Beitragsniveau ×', BOLD); r += 1
    sz_start = r
    for key, label in (('pessimistisch', '1 pessimistisch'), ('basis', '2 Basis'), ('optimistisch', '3 optimistisch')):
        s = a['szenarien'][key]
        put(an, f'B{r}', label); put(an, f'C{r}', s['neuzugaenge_faktor'], BLUE, '0.00'); put(an, f'D{r}', s['kuendigungsquote_monat'], BLUE, PCT2); put(an, f'E{r}', s['beitrag_faktor'], BLUE, '0.00')
        r += 1
    names.add_range('Sz_Neu', an, f'C{sz_start}:C{r - 1}')
    names.add_range('Sz_Churn', an, f'D{sz_start}:D{r - 1}')
    names.add_range('Sz_Beitrag', an, f'E{sz_start}:E{r - 1}')
    inp('Neu_Faktor', 'Aktiver Faktor Neuzugänge', '=INDEX(Sz_Neu,Szenario)', '0.00', '', 'Formel (Szenario)', formula=True)
    inp('Churn', 'Aktive Kündigungsquote je Monat', '=INDEX(Sz_Churn,Szenario)', PCT2, '', 'Formel (Szenario) – 2,5 %/Monat ≈ 26 % p. a.', formula=True)
    inp('Beitrag_Faktor', 'Aktiver Faktor Beitragsniveau', '=INDEX(Sz_Beitrag,Szenario)', '0.00', '', 'Formel (Szenario)', formula=True)

    head('Personal (Stellenplan, Bruttogehalt je Monat und Person)')
    inp('AG_Anteil', 'Arbeitgeberanteil Sozialversicherung (sv-pflichtig)', pe['ag_anteil_sozialversicherung'], PCT, '', 'inkl. Umlagen')
    inp('Mini_Abg', 'Pauschalabgaben Minijob', pe['minijob_pauschalabgaben'], PCT, '', '')
    inp('Steig', 'Gehaltssteigerung je Jahr', pe['gehaltssteigerung_prozent_pa'] / 100, PCT, '% p. a.', 'ab Jahr 2')
    put(an, f'B{r}', 'Rolle', BOLD); put(an, f'C{r}', 'Anzahl', BOLD); put(an, f'D{r}', 'Brutto/Monat', BOLD); put(an, f'E{r}', 'Minijob (1/0) | ab Monat', BOLD); r += 1
    # Stellen: Spalten C Anzahl, D Brutto, F Minijob, G ab Monat (E bleibt Beschreibung)
    setw(an, {'F': 10, 'G': 10})
    p_start = r
    for s in pe['stellen']:
        put(an, f'B{r}', s['rolle']); put(an, f'C{r}', s['anzahl'], BLUE, NUM); put(an, f'D{r}', s['brutto_monat'], BLUE, EUR)
        put(an, f'F{r}', 1 if s['minijob'] else 0, BLUE, '0'); put(an, f'G{r}', s['ab_monat'], BLUE, '0')
        put(an, f'E{r}', 'Minijob' if s['minijob'] else 'sozialversicherungspflichtig', GREY)
        r += 1
    put(an, f'F{p_start - 1}', 'Minijob', BOLD); put(an, f'G{p_start - 1}', 'ab Monat', BOLD)
    names.add_range('Pers_Anzahl', an, f'C{p_start}:C{r - 1}')
    names.add_range('Pers_Brutto', an, f'D{p_start}:D{r - 1}')
    names.add_range('Pers_Mini', an, f'F{p_start}:F{r - 1}')
    names.add_range('Pers_Ab', an, f'G{p_start}:G{r - 1}')
    inp('Personal_Voll', 'Personalkosten je Monat bei voller Besetzung (AG-Brutto, Jahr 1)', '=SUMPRODUCT(Pers_Anzahl*Pers_Brutto*(1+Pers_Mini*Mini_Abg+(1-Pers_Mini)*AG_Anteil))', EUR, '€/Monat', 'Formel', formula=True)
    inp('VZAE', 'Vollzeitäquivalente (ohne Minijobs, Teilzeit = 0,5)', '=SUMPRODUCT(Pers_Anzahl*(1-Pers_Mini)*((Pers_Brutto>=2000)+(Pers_Brutto<2000)*0.5))', '0.0', 'VZÄ', 'Formel (Teilzeit erkannt an Brutto < 2.000 €)', formula=True)

    head('Sachkosten je Monat (netto)')
    inp('Marketing_laufend', 'Marketing laufend (Online, Social, Print, Aktionen)', sk['marketing_laufend'], EUR, '€/Monat', '')
    inp('Marketing_Pct', 'Marketing zusätzlich in % vom Umsatz', sk['marketing_prozent_vom_umsatz'], PCT, '', '')
    inp('Marketing_Eroeff', 'Marketing Eröffnungsphase zusätzlich (Monate 1–3)', sk['marketing_eroeffnung_monate_1_3'], EUR, '€/Monat', '')
    inp('Marketing_Vorverkauf', 'Marketing Vorverkauf (gesamt, vor Eröffnung)', vo['marketing_vorverkauf_gesamt'], EUR, '€', 'verteilt auf die letzten Vorlaufmonate')
    inp('Marketing_Vorverkauf_Monate', 'Vorverkaufsdauer', vo['marketing_vorverkauf_monate'], NUM, 'Monate', '')
    inp('Versicherung', 'Versicherungen (Betriebshaftpflicht, Inhalt, Betriebsunterbrechung, Rechtsschutz)', sk['versicherungen'], EUR, '€/Monat', 'ohne Vorsteuer (Versicherungsteuer)')
    inp('GEMA', 'GEMA / GVL (Musik in Trainingsfläche und Kursen)', sk['gema_gvl'], EUR, '€/Monat', '')
    inp('Software', 'Studiosoftware, Zutritt, App, Lastschriftmanagement', sk['studiosoftware_zutritt'], EUR, '€/Monat', '')
    inp('Payment_Pct', 'Zahlungsverkehr (Lastschrift, Karten, Rücklastschriften) in % vom Beitragsumsatz brutto', sk['payment_prozent_vom_beitrag_brutto'], PCT2, '', '')
    inp('StB', 'Steuerberatung, Buchhaltung, Lohnabrechnung, Jahresabschluss', sk['steuerberater_buchhaltung'], EUR, '€/Monat', 'ab Monat −4')
    inp('Reinigung', 'Reinigungsmaterial, Hygiene, Handtuchservice, Verbrauchsmaterial', sk['reinigung_material_hygiene'], EUR, '€/Monat', 'Reinigungspersonal im Stellenplan')
    inp('Telko', 'Telekommunikation, Internet, IT, Musik-Streaming', sk['telekommunikation_it'], EUR, '€/Monat', '')
    inp('Wartung_Pct', 'Wartung/Instandhaltung Geräte in % vom Gerätewert je Jahr', sk['wartung_prozent_geraetewert_pa'], PCT, '% p. a.', 'Cardio-Wartungsverträge, Ersatzteile, Polster')
    inp('Instand_m2', 'Instandhaltung Gebäude/Technik', sk['instandhaltung_gebaeude_eur_m2_pa'], EUR2, '€/m²/Jahr', 'Lüftung, Sauna, Sanitär')
    inp('Sonstiges', 'Sonstiges (Büro, Fortbildung, Beiträge IHK/DSSV, Kfz, Kleinreparaturen)', sk['sonstiges'], EUR, '€/Monat', '')
    inp('Gruendung', 'Gründungskosten GmbH (Notar, Register, Anmeldungen)', vo['gruendungskosten_gmbh'], EUR, '€', 'Investition/Anlaufkosten')
    inp('Beratung', 'Beratung Businessplan / Steuerberatung Gründung', vo['beratung_businessplan_steuer'], EUR, '€', 'Investition/Anlaufkosten')

    head('Investition und Abschreibung')
    inp('Unvorh_Pct', 'Unvorhergesehenes in % der Investitionssumme', inv_a['unvorhergesehen_prozent'], PCT, '', 'auf Geräte, Ausbau, Sanitär, Planung, Sonstiges')
    inp('AfA_Geraete', 'Nutzungsdauer Geräte', inv_a['afa_jahre']['geraete'], NUM, 'Jahre', 'AfA-Tabelle: Fitnessgeräte')
    inp('AfA_Ausbau', 'Nutzungsdauer Mieterausbau', inv_a['afa_jahre']['ausbau'], NUM, 'Jahre', 'Mietvertragslaufzeit')
    inp('AfA_Sonst', 'Nutzungsdauer Sonstiges (Software, Zutritt)', inv_a['afa_jahre']['sonstiges'], NUM, 'Jahre', '')

    head('Finanzierung')
    inp('EK', 'Eigenkapital (Bareinlage Gründer/Familie, Stammkapital + Rücklage)', fin['eigenkapital'], EUR, '€', 'Planannahme – bitte mit tatsächlich verfügbaren Mitteln füllen', yellow=True)
    nd = len(fin['darlehen'])
    for i, d in enumerate(fin['darlehen'], start=1):
        inp(f'Darl{i}_Name', f'Darlehen {i}', d['name'], '@', '', '')
        an[f'C{r - 1}'].alignment = Alignment(wrap_text=True)
        if d.get('plug'):
            others = '-'.join(f'Darl{k}_Betrag' for k in range(1, nd + 1) if k != i)
            inp(f'Darl{i}_Betrag', f'Darlehen {i}: Betrag (Restfinanzierung = Kapitalbedarf − Eigenkapital − übrige Darlehen)', f'=MAX(0,Kapitalbedarf-EK-{others})', EUR, '€', 'Formel; bei Bedarf durch festen Betrag ersetzen', formula=True)
        else:
            inp(f'Darl{i}_Betrag', f'Darlehen {i}: Betrag', d['betrag'], EUR, '€', '', yellow=True)
        inp(f'Darl{i}_EK', f'Darlehen {i}: zählt als wirtschaftliches Eigenkapital (1 = ja, 0 = nein)', 1 if d.get('eigenkapitalaehnlich') else 0, '0', '', 'nachrangige stille Beteiligung (MBG) = 1')
        inp(f'Darl{i}_Zins', f'Darlehen {i}: Zinssatz', d['zins_pa'], PCT2, '% p. a.', 'Planannahme (Programmzins, Preisklasse Hausbank)')
        inp(f'Darl{i}_Laufzeit', f'Darlehen {i}: Laufzeit', d['laufzeit_jahre'], NUM, 'Jahre', '')
        inp(f'Darl{i}_Frei', f'Darlehen {i}: tilgungsfreie Anlaufjahre', d['tilgungsfreie_jahre'], NUM, 'Jahre', 'danach gleiche Tilgungsraten (linear)')
        inp(f'Darl{i}_Auszahlung', f'Darlehen {i}: Auszahlung im Monat', d['auszahlung_monat'], '0', 'Monat', 'negativ = Vorlaufmonat')
        inp(f'Darl{i}_Haftung', f'Darlehen {i}: Haftungsfreistellung / Bürgschaftsquote', d['haftungsfreistellung'], PCT, '', 'KfW-Haftungsfreistellung bzw. Bürgschaftsbank')
    inp('Reserve', 'Anlaufreserve / Betriebsmittel (Deckung der ersten 24 Monate)', fin['betriebsmittelreserve'], EUR, '€', 'Blatt Kennzahlen: berechneter Bedarf inkl. Puffer – Eingabe sollte ≥ Bedarf sein', yellow=True)
    inp('Puffer_Pct', 'Sicherheitspuffer auf den berechneten Reservebedarf', fin['reserve_sicherheitspuffer_prozent'], PCT, '', '')
    inp('KK_Rahmen', 'Kontokorrentrahmen (für USt-Vorfinanzierung, nicht im Kapitalbedarf)', fin['kontokorrent_rahmen'], EUR, '€', 'Vorsteuer auf Investitionen wird vom Finanzamt erstattet')

    head('Steuern')
    inp('Hebesatz', 'Gewerbesteuer-Hebesatz Standortgemeinde', st['gewerbesteuer_hebesatz'], '0%', '', 'Planannahme, gemeindeabhängig')
    inp('Messzahl', 'Gewerbesteuer-Messzahl', st['gewerbesteuer_messzahl'], PCT, '', '')
    inp('KSt', 'Körperschaftsteuer', st['koerperschaftsteuer'], PCT, '', '')
    inp('Soli', 'Solidaritätszuschlag auf KSt', st['soli'], PCT, '', '')
    inp('Steuersatz', 'Ertragsteuersatz gesamt (GmbH)', '=Messzahl*Hebesatz+KSt*(1+Soli)', PCT2, '', 'Formel', formula=True)
    inp('UST_Verz', 'Verzögerung USt-Zahlung/-Erstattung', st['ust_erstattung_verzoegerung_monate'], '0', 'Monate', '0 oder 1 (Voranmeldung im Folgemonat)')
    an.freeze_panes = 'B4'

    # ======================= Investition =======================
    iv = wb.create_sheet('Investition')
    setw(iv, {'A': 3, 'B': 14, 'C': 58, 'D': 14, 'E': 9, 'F': 9, 'G': 9, 'H': 8, 'I': 12, 'J': 46})
    put(iv, 'B1', 'Investition und Kapitalbedarf (netto, EUR)', H1)
    put(iv, 'B2', 'Gerätepositionen und Ausbau aus der GymPlanner-Kalkulation (Blatt-Quelle: Stückliste No.1, Miete 8,50 €/m²). Vor Vertragsabschluss durch Angebote ersetzen. Zahlungsplan rechts: Verteilung auf die Vorlaufmonate.', GREY)
    hdr = ['Gruppe', 'Position', 'Betrag netto', 'AfA Jahre', 'AfA je Jahr', 'Vorsteuer', 'ab Monat', 'Monate', 'Hinweis']
    for i, h in enumerate(hdr):
        put(iv, f'{L(2 + i)}4', h, BOLD, fill=HEAD)
    inv = M.investition(a, daten)
    r = 5
    inv_first = r
    group_label = {'geraete': 'Geräte', 'import': 'Geräte', 'ausbau': 'Ausbau', 'sanitaer': 'Sanitär', 'planung': 'Planung', 'sonstiges': 'Sonstiges', 'unvorhergesehen': 'Reserve', 'kaution': 'Kaution', 'gruendung': 'Gründung'}
    afa_name = {'geraete': 'AfA_Geraete', 'import': 'AfA_Geraete', 'ausbau': 'AfA_Ausbau', 'sanitaer': 'AfA_Ausbau', 'planung': 'AfA_Ausbau', 'sonstiges': 'AfA_Sonst', 'unvorhergesehen': 'AfA_Ausbau'}
    rows_basis = []
    for p in inv:
        put(iv, f'B{r}', group_label[p.gruppe])
        put(iv, f'C{r}', p.bezeichnung)
        if p.gruppe == 'unvorhergesehen':
            put(iv, f'D{r}', f'=ROUND(Unvorh_Pct*SUM(D{inv_first}:D{r - 1}),0)', BLACK, EUR)
        elif p.gruppe == 'kaution':
            put(iv, f'D{r}', '=ROUND(Kaution_Monate*(Miete_kalt+BK+Energie)*Flaeche,0)', BLACK, EUR)
        elif p.gruppe == 'gruendung':
            put(iv, f'D{r}', '=Gruendung' if 'Gründungskosten' in p.bezeichnung else '=Beratung', GREEN, EUR)
        else:
            put(iv, f'D{r}', round(p.betrag), BLUE, EUR)
            rows_basis.append(r)
        if p.afa_jahre:
            put(iv, f'E{r}', f'={afa_name[p.gruppe]}', GREEN, '0')
            put(iv, f'F{r}', f'=IF(E{r}>0,D{r}/E{r},0)', BLACK, EUR)
        else:
            put(iv, f'E{r}', 0, BLUE, '0'); put(iv, f'F{r}', f'=IF(E{r}>0,D{r}/E{r},0)', BLACK, EUR)
        put(iv, f'G{r}', 1 if p.vorsteuer else 0, BLUE, '0')
        start = min(m for m, _ in p.zahlung)
        put(iv, f'H{r}', start, BLUE, '0')
        put(iv, f'I{r}', len(p.zahlung), BLUE, '0')
        put(iv, f'J{r}', p.hinweis, GREY)
        r += 1
    inv_last = r - 1
    put(iv, f'C{r}', 'Investitionssumme (ohne Anlaufreserve)', BOLD); put(iv, f'D{r}', f'=SUM(D{inv_first}:D{inv_last})', BOLD, EUR); iv[f'D{r}'].border = TOP
    names.add('Inv_Summe', iv, f'D{r}')
    r += 1
    put(iv, f'C{r}', 'Anlaufreserve / Betriebsmittel (24 Monate)', BOLD); put(iv, f'D{r}', '=Reserve', GREEN, EUR)
    r += 1
    put(iv, f'C{r}', 'Kapitalbedarf gesamt', BOLD); put(iv, f'D{r}', f'=Inv_Summe+Reserve', BOLD, EUR); iv[f'D{r}'].border = TOP
    names.add('Kapitalbedarf', iv, f'D{r}')
    r += 2
    put(iv, f'C{r}', 'Gerätewert (Geräte + Import) – Basis Wartung', BOLD); put(iv, f'D{r}', f'=SUMIF(B{inv_first}:B{inv_last},"Geräte",D{inv_first}:D{inv_last})', BLACK, EUR)
    names.add('Geraetewert', iv, f'D{r}')
    r += 1
    put(iv, f'C{r}', 'Gründungs-/Beratungskosten (Aufwand Jahr 1)', BOLD); put(iv, f'D{r}', f'=SUMIF(B{inv_first}:B{inv_last},"Gründung",D{inv_first}:D{inv_last})', BLACK, EUR)
    names.add('Aufwand_Gruendung', iv, f'D{r}')
    r += 1
    put(iv, f'C{r}', 'Umsatzsteuer auf Investitionen (Vorsteuer, wird erstattet – Zwischenfinanzierung)', BOLD); put(iv, f'D{r}', f'=SUMPRODUCT(D{inv_first}:D{inv_last},G{inv_first}:G{inv_last})*UST', BLACK, EUR)
    names.add('Inv_Vorsteuer_Gesamt', iv, f'D{r}')
    r += 2
    # Zusammenfassung je Gruppe
    put(iv, f'C{r}', 'Zusammenfassung je Gruppe', H2); r += 1
    for g in ['Geräte', 'Ausbau', 'Sanitär', 'Planung', 'Sonstiges', 'Reserve', 'Kaution', 'Gründung']:
        put(iv, f'C{r}', g); put(iv, f'D{r}', f'=SUMIF($B${inv_first}:$B${inv_last},"{g}",$D${inv_first}:$D${inv_last})', BLACK, EUR); r += 1
    names.add_range('Inv_Gruppe', iv, f'B{inv_first}:B{inv_last}')
    names.add_range('Inv_Betrag', iv, f'D{inv_first}:D{inv_last}')
    names.add_range('Inv_AfaJahre', iv, f'E{inv_first}:E{inv_last}')
    names.add_range('Inv_AfaJeJahr', iv, f'F{inv_first}:F{inv_last}')
    names.add_range('Inv_Vst', iv, f'G{inv_first}:G{inv_last}')
    names.add_range('Inv_Start', iv, f'H{inv_first}:H{inv_last}')
    names.add_range('Inv_Monate', iv, f'I{inv_first}:I{inv_last}')
    # Zahlungsplan: Spalten L.. für Monate -V..-1 und 1..3 (Geräte könnten auch in Monat 1 fällig sein)
    zp_months = list(range(-V, 0)) + [1, 2, 3]
    put(iv, 'L3', 'Zahlungsplan (netto) je Monat', H2)
    put(iv, 'L4', 'Monat', BOLD, fill=HEAD)
    for j, m in enumerate(zp_months):
        col = L(13 + j)
        put(iv, f'{col}4', m, BOLD, '0', fill=HEAD)
        iv.column_dimensions[col].width = 11
        for rr in range(inv_first, inv_last + 1):
            put(iv, f'{col}{rr}', f'=IF(AND({col}$4>=$H{rr},{col}$4<$H{rr}+$I{rr}),$D{rr}/$I{rr},0)', BLACK, EUR)
        put(iv, f'{col}{inv_last + 1}', f'=SUM({col}{inv_first}:{col}{inv_last})', BOLD, EUR)
        put(iv, f'{col}{inv_last + 2}', f'=SUMPRODUCT({col}{inv_first}:{col}{inv_last},$G${inv_first}:$G${inv_last})*UST', BLACK, EUR)
    put(iv, f'L{inv_last + 1}', 'Summe netto', BOLD)
    put(iv, f'L{inv_last + 2}', 'Vorsteuer', BOLD)
    names.add_range('ZP_Monat', iv, f'M4:{L(13 + len(zp_months) - 1)}4')
    names.add_range('ZP_Netto', iv, f'M{inv_last + 1}:{L(13 + len(zp_months) - 1)}{inv_last + 1}')
    names.add_range('ZP_Vst', iv, f'M{inv_last + 2}:{L(13 + len(zp_months) - 1)}{inv_last + 2}')
    iv.freeze_panes = 'D5'

    # ======================= Zeitachse (Hilfsblatt-Zeilen auf Kosten) =======================
    # Gemeinsame Spaltenlogik: Spalte C = erster Monat (-V). Spaltenindex für Monat m:
    def col_of(m: int) -> str:
        i = monate_all.index(m)
        return L(3 + i)

    first_col = col_of(monate_all[0])
    last_col = col_of(MONATE)
    col_m1 = col_of(1)

    def month_header(ws: Worksheet, row: int, months: list[int], start_col: int = 3, with_year: bool = True) -> None:
        put(ws, f'B{row}', 'Monat', BOLD, fill=HEAD)
        for j, m in enumerate(months):
            c = L(start_col + j)
            put(ws, f'{c}{row}', m, BOLD, '0', fill=HEAD)
            ws.column_dimensions[c].width = 11
        if with_year:
            put(ws, f'B{row + 1}', 'Jahr', BOLD)
            put(ws, f'B{row + 2}', 'Periode', BOLD)
            for j, m in enumerate(months):
                c = L(start_col + j)
                put(ws, f'{c}{row + 1}', f'=IF({c}{row}<1,1,INT(({c}{row}-1)/12)+1)', BLACK, '0')
                put(ws, f'{c}{row + 2}', j + 1, BLACK, '0')

    # ======================= Mitglieder =======================
    mg_ws = wb.create_sheet('Mitglieder')
    setw(mg_ws, {'A': 3, 'B': 44})
    put(mg_ws, 'B1', 'Mitgliederentwicklung je Monat', H1)
    put(mg_ws, 'B2', 'Bestand Ende = MIN(Obergrenze; Anfang + Neuzugänge − Kündigungen). Kündigungen = Anfangsbestand × Kündigungsquote. Monat 1 startet mit dem Vorverkaufsbestand.', GREY)
    months_op = list(range(1, MONATE + 1))
    month_header(mg_ws, 4, months_op)
    rows_mg = {'anfang': 7, 'neu': 8, 'kuend': 9, 'ende': 10, 'avg': 11, 'spitze': 12, 'auslastung': 13}
    put(mg_ws, 'B7', 'Bestand Anfang', BLACK); put(mg_ws, 'B8', 'Neuzugänge', BLACK); put(mg_ws, 'B9', 'Kündigungen', BLACK); put(mg_ws, 'B10', 'Bestand Ende', BOLD)
    put(mg_ws, 'B11', 'Durchschnittsbestand (Abrechnungsbasis)', BLACK); put(mg_ws, 'B12', 'Gleichzeitig anwesend zur Spitzenzeit', BLACK); put(mg_ws, 'B13', 'Auslastung Spitzenzeit (gegen Kapazität)', BLACK)
    for j, m in enumerate(months_op):
        c = L(3 + j)
        prev = L(2 + j)
        put(mg_ws, f'{c}7', '=Vorverkauf*Neu_Faktor' if m == 1 else f'={prev}10', BLACK, NUM)
        put(mg_ws, f'{c}8', f'=Neu_Faktor*IF({c}$4<=3,Neu_1,IF({c}$4<=12,Neu_2,IF({c}$4<=24,Neu_3,IF({c}$4<=36,Neu_4,Neu_5))))', BLACK, NUM)
        put(mg_ws, f'{c}9', f'={c}7*Churn', BLACK, NUM)
        put(mg_ws, f'{c}10', f'=MIN(Kapazitaet_Max,{c}7+{c}8-{c}9)', BOLD, NUM)
        put(mg_ws, f'{c}11', f'=({c}7+{c}10)/2', BLACK, NUM)
        put(mg_ws, f'{c}12', f'={c}10*Spitzenlast', BLACK, NUM)
        put(mg_ws, f'{c}13', f'={c}12/Personen', BLACK, PCT)
    names.add_range('Mg_Monat', mg_ws, f'C4:{L(2 + MONATE)}4')
    names.add_range('Mg_Jahr', mg_ws, f'C5:{L(2 + MONATE)}5')
    names.add_range('Mg_Neu', mg_ws, f'C8:{L(2 + MONATE)}8')
    names.add_range('Mg_Ende', mg_ws, f'C10:{L(2 + MONATE)}10')
    names.add_range('Mg_Avg', mg_ws, f'C11:{L(2 + MONATE)}11')
    # Jahreswerte
    put(mg_ws, 'B16', 'Jahreswerte', H2)
    put(mg_ws, 'B17', 'Jahr', BOLD, fill=HEAD)
    for j in range(1, 6):
        c = L(2 + j)
        put(mg_ws, f'{c}17', j, BOLD, '0', fill=HEAD)
        put(mg_ws, f'{c}18', f'=SUMIF(Mg_Jahr,{c}17,Mg_Neu)', BLACK, NUM)
        put(mg_ws, f'{c}19', f'=INDEX(Mg_Ende,{c}17*12)', BLACK, NUM)
        put(mg_ws, f'{c}20', f'={c}19/Flaeche', BLACK, '0.00')
    put(mg_ws, 'B18', 'Neuzugänge im Jahr'); put(mg_ws, 'B19', 'Bestand am Jahresende'); put(mg_ws, 'B20', 'Mitglieder je m² Bruttofläche')
    mg_ws.freeze_panes = 'C7'

    # ======================= Umsatz =======================
    um = wb.create_sheet('Umsatz')
    setw(um, {'A': 3, 'B': 48})
    put(um, 'B1', 'Umsatzplanung je Monat (netto, EUR)', H1)
    put(um, 'B2', 'Beiträge = Durchschnittsbestand × Ø Beitrag brutto ÷ (1 + USt). Beitragsanpassung ab Jahr 3. Nebenumsätze je Mitglied und Monat.', GREY)
    month_header(um, 4, months_op)
    labels_um = ['Durchschnittsbestand', 'Ø Beitrag brutto (inkl. Anpassung)', 'Mitgliedsbeiträge netto', 'Servicepauschale netto', 'Aufnahmegebühren netto', 'Nebenumsatz Getränke/Shakes netto', 'Nebenumsatz Wellness/Solarium netto', 'Nebenumsatz PT/Kurse netto', 'Nebenumsatz Merchandise netto', 'Umsatz gesamt netto', 'Umsatzsteuer auf Umsatz', 'Wareneinsatz Getränke/Shakes']
    for i, lab in enumerate(labels_um):
        put(um, f'B{7 + i}', lab, BOLD if lab.startswith('Umsatz gesamt') else BLACK)
    for j, m in enumerate(months_op):
        c = L(3 + j)
        put(um, f'{c}7', f"=Mitglieder!{c}11", GREEN, NUM)
        put(um, f'{c}8', f'=ARPU_brutto*(1+Beitrag_Erhoehung)^MAX(0,{c}$5-2)', BLACK, EUR2)
        put(um, f'{c}9', f'={c}7*{c}8/(1+UST)', BLACK, EUR)
        put(um, f'{c}10', f'={c}7*Service_Halbjahr/6/(1+UST)', BLACK, EUR)
        put(um, f'{c}11', f"=Mitglieder!{c}8*Aufnahme_Anteil*Aufnahme/(1+UST)", BLACK, EUR)
        put(um, f'{c}12', f'={c}7*Neben_Getraenke/(1+UST)', BLACK, EUR)
        put(um, f'{c}13', f'={c}7*Neben_Wellness/(1+UST)', BLACK, EUR)
        put(um, f'{c}14', f'={c}7*Neben_PT/(1+UST)', BLACK, EUR)
        put(um, f'{c}15', f'={c}7*Neben_Merch/(1+UST)', BLACK, EUR)
        put(um, f'{c}16', f'=SUM({c}9:{c}15)', BOLD, EUR)
        put(um, f'{c}17', f'={c}16*UST', BLACK, EUR)
        put(um, f'{c}18', f'={c}12*Neben_WE', BLACK, EUR)
    for nm, row in (('Um_Monat', 4), ('Um_Jahr', 5), ('Um_Beitrag', 9), ('Um_Service', 10), ('Um_Aufnahme', 11), ('Um_Gesamt', 16), ('Um_USt', 17), ('Um_WE', 18)):
        names.add_range(nm, um, f'C{row}:{L(2 + MONATE)}{row}')
    put(um, 'B21', 'Jahreswerte (netto)', H2)
    put(um, 'B22', 'Jahr', BOLD, fill=HEAD)
    for j in range(1, 6):
        c = L(2 + j)
        put(um, f'{c}22', j, BOLD, '0', fill=HEAD)
        put(um, f'{c}23', f'=SUMIF(Um_Jahr,{c}22,Um_Beitrag)+SUMIF(Um_Jahr,{c}22,Um_Service)', BLACK, EUR)
        put(um, f'{c}24', f'=SUMIF(Um_Jahr,{c}22,Um_Aufnahme)', BLACK, EUR)
        put(um, f'{c}25', f'=SUMIF(Um_Jahr,{c}22,Um_Gesamt)-{c}23-{c}24', BLACK, EUR)
        put(um, f'{c}26', f'=SUMIF(Um_Jahr,{c}22,Um_Gesamt)', BOLD, EUR)
        put(um, f'{c}27', f'={c}26/Flaeche', BLACK, EUR)
    put(um, 'B23', 'Beiträge inkl. Servicepauschale'); put(um, 'B24', 'Aufnahmegebühren'); put(um, 'B25', 'Nebenumsätze'); put(um, 'B26', 'Umsatz gesamt', BOLD); put(um, 'B27', 'Umsatz je m² Bruttofläche')
    um.freeze_panes = 'C7'

    # ======================= Kosten =======================
    ko = wb.create_sheet('Kosten')
    setw(ko, {'A': 3, 'B': 52})
    put(ko, 'B1', 'Kostenplanung je Monat (netto, EUR) – Vorlauf und Betrieb', H1)
    put(ko, 'B2', 'Monate < 1 = Vorlauf (Ausbau, Vorverkauf): Miete nach mietfreier Zeit, Personal je Eintrittsmonat, Vorverkaufsmarketing; ab Monat 1 volle Betriebskosten.', GREY)
    month_header(ko, 4, monate_all)
    kost_rows = [
        ('personal', 'Personal (AG-Brutto inkl. Sozialabgaben)', 0),
        ('miete', 'Kaltmiete', 1), ('bk', 'Betriebskosten', 1), ('energie', 'Energie und Wasser', 1),
        ('marketing', 'Marketing (laufend, Eröffnung, Vorverkauf)', 1), ('versicherung', 'Versicherungen', 0), ('gema', 'GEMA / GVL', 1),
        ('software', 'Studiosoftware / Zutritt', 1), ('payment', 'Zahlungsverkehr', 0), ('stb', 'Steuerberatung / Buchhaltung', 1),
        ('reinigung', 'Reinigung / Hygiene / Verbrauch', 1), ('telko', 'Telekommunikation / IT', 1), ('wartung', 'Wartung Geräte', 1),
        ('instand', 'Instandhaltung Gebäude/Technik', 1), ('sonstiges', 'Sonstiges', 1), ('we', 'Wareneinsatz Getränke/Shakes', 1),
    ]
    row_of = {k: 7 + i for i, (k, _, _) in enumerate(kost_rows)}
    for k, lab, _ in kost_rows:
        put(ko, f'B{row_of[k]}', lab)
    r_sum = 7 + len(kost_rows)
    put(ko, f'B{r_sum}', 'Kosten gesamt (netto)', BOLD)
    put(ko, f'B{r_sum + 1}', 'davon vorsteuerfähig', BLACK)
    put(ko, f'B{r_sum + 2}', 'Vorsteuer auf Kosten', BLACK)
    put(ko, f'B{r_sum + 3}', 'Vorsteuer-Flag je Zeile → Spalte C', GREY)
    vst_flag_col = 'C'
    # Vorsteuer-Flags in eine eigene Spalte? Spalte C ist Monat -V. Wir legen die Flags in Spalte BQ (nach den Monaten).
    flag_col = L(3 + len(monate_all) + 1)
    put(ko, f'{flag_col}4', 'Vorsteuer (1/0)', BOLD, fill=HEAD)
    ko.column_dimensions[flag_col].width = 14
    for k, _, vf in kost_rows:
        put(ko, f'{flag_col}{row_of[k]}', vf, BLUE, '0')
    put(ko, f'B{r_sum + 3}', f'Vorsteuer-Flag je Zeile in Spalte {flag_col} (1 = Vorsteuer abziehbar)', GREY)
    for j, m in enumerate(monate_all):
        c = L(3 + j)
        mref = f'{c}$4'
        jref = f'{c}$5'
        um_col = col_of(m) if m >= 1 else None
        # Umsatzblatt hat nur Monate 1..60 ab Spalte C → Spalte im Umsatzblatt = L(3 + (m-1))
        uc = L(3 + (m - 1)) if m >= 1 else None
        put(ko, f'{c}{row_of["personal"]}', f'=SUMPRODUCT((Pers_Ab<={mref})*Pers_Anzahl*Pers_Brutto*(1+Pers_Mini*Mini_Abg+(1-Pers_Mini)*AG_Anteil))*(1+Steig)^MAX(0,{jref}-1)', BLACK, EUR)
        put(ko, f'{c}{row_of["miete"]}', f'=IF(AND({mref}<1,{mref}+Vorlauf<Mietfrei),0,Flaeche*Miete_kalt*(1+Index_pa)^MAX(0,{jref}-1))', BLACK, EUR)
        put(ko, f'{c}{row_of["bk"]}', f'=IF(AND({mref}<1,{mref}+Vorlauf<Mietfrei),0,Flaeche*BK)', BLACK, EUR)
        put(ko, f'{c}{row_of["energie"]}', f'=IF({mref}<1,0,Flaeche*Energie)', BLACK, EUR)
        if m < 1:
            put(ko, f'{c}{row_of["marketing"]}', f'=IF({mref}+Vorlauf>=Vorlauf-Marketing_Vorverkauf_Monate,Marketing_Vorverkauf/Marketing_Vorverkauf_Monate,0)', BLACK, EUR)
            put(ko, f'{c}{row_of["payment"]}', 0, BLACK, EUR)
            put(ko, f'{c}{row_of["we"]}', 0, BLACK, EUR)
        else:
            put(ko, f'{c}{row_of["marketing"]}', f'=Marketing_laufend+IF({mref}<=3,Marketing_Eroeff,0)+Marketing_Pct*Umsatz!{uc}16', BLACK, EUR)
            put(ko, f'{c}{row_of["payment"]}', f'=Payment_Pct*(Umsatz!{uc}9+Umsatz!{uc}10+Umsatz!{uc}11)*(1+UST)', BLACK, EUR)
            put(ko, f'{c}{row_of["we"]}', f'=Umsatz!{uc}18', GREEN, EUR)
        put(ko, f'{c}{row_of["versicherung"]}', f'=IF({mref}<1,0,Versicherung)', BLACK, EUR)
        put(ko, f'{c}{row_of["gema"]}', f'=IF({mref}<1,0,GEMA)', BLACK, EUR)
        put(ko, f'{c}{row_of["software"]}', f'=IF({mref}<1,0,Software)', BLACK, EUR)
        put(ko, f'{c}{row_of["stb"]}', f'=IF({mref}+Vorlauf>=1,StB,0)', BLACK, EUR)
        put(ko, f'{c}{row_of["reinigung"]}', f'=IF({mref}<1,0,Reinigung)', BLACK, EUR)
        put(ko, f'{c}{row_of["telko"]}', f'=IF({mref}<1,0,Telko)', BLACK, EUR)
        put(ko, f'{c}{row_of["wartung"]}', f'=IF({mref}<1,0,Geraetewert*Wartung_Pct/12)', BLACK, EUR)
        put(ko, f'{c}{row_of["instand"]}', f'=IF({mref}<1,0,Flaeche*Instand_m2/12)', BLACK, EUR)
        put(ko, f'{c}{row_of["sonstiges"]}', f'=IF({mref}<1,0,Sonstiges)', BLACK, EUR)
        put(ko, f'{c}{r_sum}', f'=SUM({c}7:{c}{r_sum - 1})', BOLD, EUR)
        put(ko, f'{c}{r_sum + 1}', f'=SUMPRODUCT({c}7:{c}{r_sum - 1},${flag_col}$7:${flag_col}${r_sum - 1})-IF(UST_Miete=0,{c}{row_of["miete"]}+{c}{row_of["bk"]},0)', BLACK, EUR)
        put(ko, f'{c}{r_sum + 2}', f'={c}{r_sum + 1}*UST', BLACK, EUR)
    for nm, row in (('Ko_Monat', 4), ('Ko_Jahr', 5), ('Ko_Periode', 6), ('Ko_Personal', row_of['personal']), ('Ko_Miete', row_of['miete']), ('Ko_BK', row_of['bk']), ('Ko_Energie', row_of['energie']), ('Ko_Marketing', row_of['marketing']), ('Ko_Instand', row_of['instand']), ('Ko_WE', row_of['we']), ('Ko_Gesamt', r_sum), ('Ko_Vorsteuer', r_sum + 2)):
        names.add_range(nm, ko, f'{first_col}{row}:{last_col}{row}')
    # Jahreswerte
    r_j = r_sum + 6
    put(ko, f'B{r_j - 1}', 'Jahreswerte (netto; Jahr 1 inkl. Vorlauf)', H2)
    put(ko, f'B{r_j}', 'Jahr', BOLD, fill=HEAD)
    for j in range(1, 6):
        c = L(2 + j)
        put(ko, f'{c}{r_j}', j, BOLD, '0', fill=HEAD)
        for i, (k, lab, _) in enumerate(kost_rows):
            put(ko, f'{c}{r_j + 1 + i}', f'=SUMIF(Ko_Jahr,{c}${r_j},${first_col}{row_of[k]}:${last_col}{row_of[k]})', BLACK, EUR)
        put(ko, f'{c}{r_j + 1 + len(kost_rows)}', f'=SUM({c}{r_j + 1}:{c}{r_j + len(kost_rows)})', BOLD, EUR)
    for i, (k, lab, _) in enumerate(kost_rows):
        put(ko, f'B{r_j + 1 + i}', lab)
    put(ko, f'B{r_j + 1 + len(kost_rows)}', 'Kosten gesamt', BOLD)
    ko.freeze_panes = 'C7'

    # ======================= Finanzierung =======================
    fz = wb.create_sheet('Finanzierung')
    setw(fz, {'A': 3, 'B': 50})
    put(fz, 'B1', 'Finanzierungsplan – Mittelverwendung, Mittelherkunft, Tilgungspläne', H1)
    put(fz, 'B3', 'Mittelverwendung', H2, fill=HEAD); put(fz, 'C3', 'EUR', BOLD, fill=HEAD); put(fz, 'D3', 'Anteil', BOLD, fill=HEAD)
    fz.column_dimensions['C'].width = 16; fz.column_dimensions['D'].width = 10
    mv = [('Geräte und Import', '=SUMIF(Inv_Gruppe,"Geräte",Inv_Betrag)'), ('Ausbau (Mieterausbau, Lüftung, Böden, Spiegel, Brandschutz)', '=SUMIF(Inv_Gruppe,"Ausbau",Inv_Betrag)'), ('Sanitärinstallation', '=SUMIF(Inv_Gruppe,"Sanitär",Inv_Betrag)'), ('Planung und Genehmigung', '=SUMIF(Inv_Gruppe,"Planung",Inv_Betrag)'), ('Sonstiges (Zutritt, Software, Kleinmaterial)', '=SUMIF(Inv_Gruppe,"Sonstiges",Inv_Betrag)'), ('Unvorhergesehenes', '=SUMIF(Inv_Gruppe,"Reserve",Inv_Betrag)'), ('Mietkaution', '=SUMIF(Inv_Gruppe,"Kaution",Inv_Betrag)'), ('Gründung und Beratung', '=SUMIF(Inv_Gruppe,"Gründung",Inv_Betrag)'), ('Anlaufreserve / Betriebsmittel (24 Monate)', '=Reserve')]
    for i, (lab, f) in enumerate(mv):
        put(fz, f'B{4 + i}', lab); put(fz, f'C{4 + i}', f, GREEN, EUR); put(fz, f'D{4 + i}', f'=C{4 + i}/$C${4 + len(mv)}', BLACK, PCT)
    rt = 4 + len(mv)
    put(fz, f'B{rt}', 'Kapitalbedarf gesamt', BOLD); put(fz, f'C{rt}', f'=SUM(C4:C{rt - 1})', BOLD, EUR); fz[f'C{rt}'].border = TOP; put(fz, f'D{rt}', f'=C{rt}/C{rt}', BLACK, PCT)
    put(fz, f'B{rt + 1}', 'Kontrolle: = Blatt Investition', GREY); put(fz, f'C{rt + 1}', f'=C{rt}-Kapitalbedarf', GREY, EUR)
    rh = rt + 3
    put(fz, f'B{rh}', 'Mittelherkunft', H2, fill=HEAD); put(fz, f'C{rh}', 'EUR', BOLD, fill=HEAD); put(fz, f'D{rh}', 'Anteil', BOLD, fill=HEAD)
    put(fz, f'B{rh + 1}', 'Eigenkapital (Bareinlage)'); put(fz, f'C{rh + 1}', '=EK', GREEN, EUR)
    nd = len(fin['darlehen'])
    for i in range(1, nd + 1):
        put(fz, f'B{rh + 1 + i}', f'=Darl{i}_Name', GREEN); put(fz, f'C{rh + 1 + i}', f'=Darl{i}_Betrag', GREEN, EUR)
        fz[f'B{rh + 1 + i}'].alignment = Alignment(wrap_text=True)
    rs = rh + 2 + nd
    put(fz, f'B{rs}', 'Summe Mittelherkunft', BOLD); put(fz, f'C{rs}', f'=SUM(C{rh + 1}:C{rs - 1})', BOLD, EUR); fz[f'C{rs}'].border = TOP
    for i in range(rh + 1, rs + 1):
        put(fz, f'D{i}', f'=C{i}/$C${rs}', BLACK, PCT)
    put(fz, f'B{rs + 1}', 'Über-/Unterdeckung (Mittelherkunft − Kapitalbedarf)', BOLD); put(fz, f'C{rs + 1}', f'=C{rs}-C{rt}', BOLD, EUR)
    names.add('Deckung', fz, f'C{rs + 1}')
    ekf = '+'.join(f'Darl{i}_Betrag*Darl{i}_EK' for i in range(1, nd + 1))
    put(fz, f'B{rs + 2}', 'Eigenkapitalquote (Bareinlage + eigenkapitalähnliche Mittel) am Kapitalbedarf'); put(fz, f'C{rs + 2}', f'=(EK+{ekf})/C{rt}', BLACK, PCT)
    names.add('EK_Quote', fz, f'C{rs + 2}')
    put(fz, f'B{rs + 3}', 'Fremdkapital gesamt (alle Darlehen und Beteiligungen)'); put(fz, f'C{rs + 3}', f'=SUM(C{rh + 2}:C{rs - 1})', BLACK, EUR)
    names.add('FK_Summe', fz, f'C{rs + 3}')
    # Tilgungspläne
    rd = rs + 6
    put(fz, f'B{rd - 1}', 'Tilgungspläne je Monat (Auszahlung im Vorlauf, Zins/Vergütung auf Restsaldo, lineare Tilgung nach tilgungsfreier Zeit; Laufzeit = tilgungsfreie Zeit → endfällig)', H2)
    month_header(fz, rd, monate_all)
    names.add_range('Fz_Monat', fz, f'{first_col}{rd}:{last_col}{rd}')
    names.add_range('Fz_Jahr', fz, f'{first_col}{rd + 1}:{last_col}{rd + 1}')
    names.add_range('Fz_Periode', fz, f'{first_col}{rd + 2}:{last_col}{rd + 2}')
    rr = rd + 4
    zins_rows, tilg_rows, ausz_rows = [], [], []
    for i in range(1, nd + 1):
        put(fz, f'B{rr}', f'=Darl{i}_Name', GREEN)
        put(fz, f'B{rr + 1}', 'Auszahlung'); put(fz, f'B{rr + 2}', 'Zinsen'); put(fz, f'B{rr + 3}', 'Tilgung'); put(fz, f'B{rr + 4}', 'Restsaldo Ende'); put(fz, f'B{rr + 5}', 'Monate seit Auszahlung')
        for j, m in enumerate(monate_all):
            c = L(3 + j)
            prev = L(2 + j)
            put(fz, f'{c}{rr + 1}', f'=IF({c}${rd}=Darl{i}_Auszahlung,Darl{i}_Betrag,0)', BLACK, EUR)
            put(fz, f'{c}{rr + 5}', f'={c}${rd + 2}-MATCH(Darl{i}_Auszahlung,Fz_Monat,0)', BLACK, '0')
            saldo_vor = f'({prev}{rr + 4}+{c}{rr + 1})' if j > 0 else f'({c}{rr + 1})'
            put(fz, f'{c}{rr + 2}', f'=IF({c}{rr + 5}>=0,{saldo_vor}*Darl{i}_Zins/12,0)', BLACK, EUR)
            put(fz, f'{c}{rr + 3}', f'=IF(Darl{i}_Laufzeit>Darl{i}_Frei,IF({c}{rr + 5}>=Darl{i}_Frei*12,MIN(Darl{i}_Betrag/((Darl{i}_Laufzeit-Darl{i}_Frei)*12),{saldo_vor}),0),IF({c}{rr + 5}=Darl{i}_Laufzeit*12-1,{saldo_vor},0))', BLACK, EUR)
            put(fz, f'{c}{rr + 4}', f'={saldo_vor}-{c}{rr + 3}', BLACK, EUR)
        zins_rows.append(rr + 2); tilg_rows.append(rr + 3); ausz_rows.append(rr + 1)
        rr += 7
    put(fz, f'B{rr}', 'Summe Auszahlungen', BOLD); put(fz, f'B{rr + 1}', 'Summe Zinsen', BOLD); put(fz, f'B{rr + 2}', 'Summe Tilgung', BOLD); put(fz, f'B{rr + 3}', 'Kapitaldienst (Zins + Tilgung)', BOLD); put(fz, f'B{rr + 4}', 'Restschuld gesamt', BOLD)
    for j, m in enumerate(monate_all):
        c = L(3 + j)
        put(fz, f'{c}{rr}', '=' + '+'.join(f'{c}{x}' for x in ausz_rows), BOLD, EUR)
        put(fz, f'{c}{rr + 1}', '=' + '+'.join(f'{c}{x}' for x in zins_rows), BOLD, EUR)
        put(fz, f'{c}{rr + 2}', '=' + '+'.join(f'{c}{x}' for x in tilg_rows), BOLD, EUR)
        put(fz, f'{c}{rr + 3}', f'={c}{rr + 1}+{c}{rr + 2}', BOLD, EUR)
        put(fz, f'{c}{rr + 4}', '=' + '+'.join(f'{c}{x + 1}' for x in tilg_rows), BOLD, EUR)
    names.add_range('Fz_Ausz', fz, f'{first_col}{rr}:{last_col}{rr}')
    names.add_range('Fz_Zins', fz, f'{first_col}{rr + 1}:{last_col}{rr + 1}')
    names.add_range('Fz_Tilg', fz, f'{first_col}{rr + 2}:{last_col}{rr + 2}')
    names.add_range('Fz_Rest', fz, f'{first_col}{rr + 4}:{last_col}{rr + 4}')
    # Kapitaldienst je Jahr (1–5) und Gesamtlaufzeit-Info
    rk = rr + 7
    put(fz, f'B{rk - 1}', 'Kapitaldienst je Jahr (Jahr 1 inkl. Vorlauf)', H2)
    put(fz, f'B{rk}', 'Jahr', BOLD, fill=HEAD); put(fz, f'B{rk + 1}', 'Zinsen'); put(fz, f'B{rk + 2}', 'Tilgung'); put(fz, f'B{rk + 3}', 'Kapitaldienst', BOLD); put(fz, f'B{rk + 4}', 'Restschuld Jahresende')
    for j in range(1, 6):
        c = L(2 + j)
        put(fz, f'{c}{rk}', j, BOLD, '0', fill=HEAD)
        put(fz, f'{c}{rk + 1}', f'=SUMIF(Fz_Jahr,{c}{rk},Fz_Zins)', BLACK, EUR)
        put(fz, f'{c}{rk + 2}', f'=SUMIF(Fz_Jahr,{c}{rk},Fz_Tilg)', BLACK, EUR)
        put(fz, f'{c}{rk + 3}', f'={c}{rk + 1}+{c}{rk + 2}', BOLD, EUR)
        put(fz, f'{c}{rk + 4}', f'=INDEX(Fz_Rest,MATCH({c}{rk}*12,Fz_Monat,0))', BLACK, EUR)
    put(fz, f'B{rk + 6}', 'Monatliche Vollrate nach der tilgungsfreien Zeit (Zins Jahr 3 + Tilgung) ≈'); put(fz, f'C{rk + 6}', f'=INDEX(Fz_Zins,MATCH(25,Fz_Monat,0))+INDEX(Fz_Tilg,MATCH(25,Fz_Monat,0))', BLACK, EUR)
    names.add('Rate_M25', fz, f'C{rk + 6}')
    fz.freeze_panes = 'C4'

    # ======================= Liquidität =======================
    lq = wb.create_sheet('Liquiditaet')
    setw(lq, {'A': 3, 'B': 52})
    put(lq, 'B1', 'Liquiditätsplan je Monat (Zahlungsströme brutto inkl. USt)', H1)
    put(lq, 'B2', 'Umsatzsteuer-Zahllast = USt auf Umsatz − Vorsteuer auf Kosten und Investition; negativer Saldo = Erstattung (Voranmeldung, Zahlung/Erstattung im Folgemonat). Steuern auf das Ergebnis werden im Folgejahr (Monat 6) fällig.', GREY)
    month_header(lq, 4, monate_all)
    lab = ['Kassenbestand Anfang', 'Einzahlungen', 'Umsatz brutto (inkl. USt)', 'USt-Erstattung vom Finanzamt', 'Eigenkapital (Einlage + Nachrang)', 'Darlehensauszahlung', 'Summe Einzahlungen', 'Auszahlungen', 'Betriebskosten brutto (Kosten + Vorsteuer)', 'Investitionen brutto (netto + Vorsteuer)', 'USt-Zahllast an das Finanzamt', 'Zinsen', 'Tilgung', 'Ertragsteuern', 'Summe Auszahlungen', 'Kassenbestand Ende', 'Hilfsrechnung', 'USt-Saldo des Monats (+ Zahllast / − Erstattung)', 'Betriebsergebnis vor Finanzierung des Kapitalbedarfs (Umsatz − Kosten − Kapitaldienst − Steuern, netto)', 'kumuliert (Basis Reservebedarf)']
    for i, l_ in enumerate(lab):
        put(lq, f'B{7 + i}', l_, BOLD if l_ in ('Einzahlungen', 'Auszahlungen', 'Summe Einzahlungen', 'Summe Auszahlungen', 'Kassenbestand Ende', 'Hilfsrechnung') else BLACK)
    R = {k: 7 + i for i, k in enumerate(['anfang', 'h_ein', 'umsatz', 'ust_erst', 'ek', 'darlehen', 'sum_ein', 'h_aus', 'kosten', 'invest', 'ust_zahl', 'zins', 'tilg', 'steuer', 'sum_aus', 'ende', 'h_hilf', 'ust_saldo', 'betrieb', 'betrieb_kum'])}
    for j, m in enumerate(monate_all):
        c = L(3 + j)
        prev = L(2 + j)
        uc = L(3 + (m - 1)) if m >= 1 else None
        put(lq, f'{c}{R["anfang"]}', 0 if j == 0 else f'={prev}{R["ende"]}', BLACK, EUR)
        put(lq, f'{c}{R["umsatz"]}', f'=Umsatz!{uc}16+Umsatz!{uc}17' if m >= 1 else 0, GREEN if m >= 1 else BLACK, EUR)
        ust_prev = f'{prev}{R["ust_saldo"]}' if j > 0 else '0'
        put(lq, f'{c}{R["ust_saldo"]}', (f'=Umsatz!{uc}17' if m >= 1 else '=0') + f'-Kosten!{c}{r_sum + 2}-IFERROR(INDEX(ZP_Vst,MATCH({c}$4,ZP_Monat,0)),0)', BLACK, EUR)
        put(lq, f'{c}{R["ust_erst"]}', f'=MAX(0,-IF(UST_Verz=0,{c}{R["ust_saldo"]},{ust_prev}))', BLACK, EUR)
        put(lq, f'{c}{R["ust_zahl"]}', f'=MAX(0,IF(UST_Verz=0,{c}{R["ust_saldo"]},{ust_prev}))', BLACK, EUR)
        put(lq, f'{c}{R["ek"]}', '=EK' if j == 0 else 0, GREEN if j == 0 else BLACK, EUR)
        put(lq, f'{c}{R["darlehen"]}', f'=Finanzierung!{c}{rr}', GREEN, EUR)
        put(lq, f'{c}{R["sum_ein"]}', f'={c}{R["umsatz"]}+{c}{R["ust_erst"]}+{c}{R["ek"]}+{c}{R["darlehen"]}', BOLD, EUR)
        put(lq, f'{c}{R["kosten"]}', f'=Kosten!{c}{r_sum}+Kosten!{c}{r_sum + 2}', GREEN, EUR)
        put(lq, f'{c}{R["invest"]}', f'=IFERROR(INDEX(ZP_Netto,MATCH({c}$4,ZP_Monat,0))+INDEX(ZP_Vst,MATCH({c}$4,ZP_Monat,0)),0)', GREEN, EUR)
        put(lq, f'{c}{R["zins"]}', f'=Finanzierung!{c}{rr + 1}', GREEN, EUR)
        put(lq, f'{c}{R["tilg"]}', f'=Finanzierung!{c}{rr + 2}', GREEN, EUR)
        put(lq, f'{c}{R["steuer"]}', f'=IF(AND({c}$4>=18,MOD({c}$4,12)=6),INDEX(GuV_Steuern,({c}$4-6)/12),0)', GREEN, EUR)
        put(lq, f'{c}{R["sum_aus"]}', f'={c}{R["kosten"]}+{c}{R["invest"]}+{c}{R["ust_zahl"]}+{c}{R["zins"]}+{c}{R["tilg"]}+{c}{R["steuer"]}', BOLD, EUR)
        put(lq, f'{c}{R["ende"]}', f'={c}{R["anfang"]}+{c}{R["sum_ein"]}-{c}{R["sum_aus"]}', BOLD, EUR)
        put(lq, f'{c}{R["betrieb"]}', (f'=Umsatz!{uc}16' if m >= 1 else '=0') + f'-Kosten!{c}{r_sum}-{c}{R["zins"]}-{c}{R["tilg"]}-{c}{R["steuer"]}', BLACK, EUR)
        put(lq, f'{c}{R["betrieb_kum"]}', f'={c}{R["betrieb"]}' if j == 0 else f'={prev}{R["betrieb_kum"]}+{c}{R["betrieb"]}', BLACK, EUR)
    names.add_range('Lq_Monat', lq, f'{first_col}4:{last_col}4')
    names.add_range('Lq_Ende', lq, f'{first_col}{R["ende"]}:{last_col}{R["ende"]}')
    names.add_range('Lq_BetriebKum', lq, f'{first_col}{R["betrieb_kum"]}:{last_col}{R["betrieb_kum"]}')
    names.add_range('Lq_Steuer', lq, f'{first_col}{R["steuer"]}:{last_col}{R["steuer"]}')
    lq.freeze_panes = 'C7'

    # ======================= GuV =======================
    gv = wb.create_sheet('GuV')
    setw(gv, {'A': 3, 'B': 56, 'C': 15, 'D': 15, 'E': 15, 'F': 15, 'G': 15})
    put(gv, 'B1', 'Rentabilitätsvorschau (Plan-GuV) Jahr 1–5, netto EUR', H1)
    put(gv, 'B2', 'Jahr 1 = Vorlaufmonate + Betriebsmonate 1–12. Abschreibung linear. Steuern mit Verlustvortrag. DSCR = (EBITDA − Steuern) ÷ Kapitaldienst.', GREY)
    put(gv, 'B4', 'Jahr', BOLD, fill=HEAD)
    glab = ['Umsatzerlöse (netto)', 'Wareneinsatz', 'Rohertrag', 'Personalkosten', 'Raumkosten (Miete, Betriebskosten, Energie, Instandhaltung)', 'Marketing', 'Sonstige betriebliche Aufwendungen', 'davon Gründungs-/Beratungskosten', 'EBITDA (Betriebsergebnis vor Abschreibung und Zinsen)', 'Abschreibungen', 'EBIT', 'Zinsaufwand', 'Ergebnis vor Steuern (EBT)', 'Verlustvortrag Anfang', 'Steuerliche Bemessungsgrundlage', 'Ertragsteuern (GewSt, KSt, Soli)', 'Jahresergebnis', 'Tilgung', 'Kapitaldienst (Zins + Tilgung)', 'Kapitaldienstfähigkeit DSCR', 'Cashflow nach Kapitaldienst (EBITDA − Steuern − Kapitaldienst)', 'Mitglieder am Jahresende', 'EBITDA-Marge', 'Personalkostenquote', 'Raumkostenquote']
    G = {k: 5 + i for i, k in enumerate(['umsatz', 'we', 'roh', 'personal', 'raum', 'marketing', 'sonst', 'gruend', 'ebitda', 'afa', 'ebit', 'zins', 'ebt', 'vv', 'basis', 'steuer', 'ergebnis', 'tilg', 'kd', 'dscr', 'cf', 'mitgl', 'marge', 'pq', 'rq'])}
    for i, l_ in enumerate(glab):
        put(gv, f'B{5 + i}', l_, BOLD if l_.startswith(('EBITDA (', 'Jahresergebnis', 'Kapitaldienstfähigkeit')) else BLACK)
    for j in range(1, 6):
        c = L(2 + j)
        prev = L(1 + j)
        put(gv, f'{c}4', j, BOLD, '0', fill=HEAD)
        put(gv, f'{c}{G["umsatz"]}', f'=SUMIF(Um_Jahr,{c}$4,Um_Gesamt)', GREEN, EUR)
        put(gv, f'{c}{G["we"]}', f'=SUMIF(Ko_Jahr,{c}$4,Ko_WE)', GREEN, EUR)
        put(gv, f'{c}{G["roh"]}', f'={c}{G["umsatz"]}-{c}{G["we"]}', BLACK, EUR)
        put(gv, f'{c}{G["personal"]}', f'=SUMIF(Ko_Jahr,{c}$4,Ko_Personal)', GREEN, EUR)
        put(gv, f'{c}{G["raum"]}', f'=SUMIF(Ko_Jahr,{c}$4,Ko_Miete)+SUMIF(Ko_Jahr,{c}$4,Ko_BK)+SUMIF(Ko_Jahr,{c}$4,Ko_Energie)+SUMIF(Ko_Jahr,{c}$4,Ko_Instand)', GREEN, EUR)
        put(gv, f'{c}{G["marketing"]}', f'=SUMIF(Ko_Jahr,{c}$4,Ko_Marketing)', GREEN, EUR)
        put(gv, f'{c}{G["gruend"]}', '=Aufwand_Gruendung' if j == 1 else 0, GREEN if j == 1 else BLACK, EUR)
        put(gv, f'{c}{G["sonst"]}', f'=SUMIF(Ko_Jahr,{c}$4,Ko_Gesamt)-{c}{G["we"]}-{c}{G["personal"]}-{c}{G["raum"]}-{c}{G["marketing"]}+{c}{G["gruend"]}', BLACK, EUR)
        put(gv, f'{c}{G["ebitda"]}', f'={c}{G["roh"]}-{c}{G["personal"]}-{c}{G["raum"]}-{c}{G["marketing"]}-{c}{G["sonst"]}', BOLD, EUR)
        put(gv, f'{c}{G["afa"]}', f'=SUMPRODUCT((Inv_AfaJahre>={c}$4)*Inv_AfaJeJahr)', GREEN, EUR)
        put(gv, f'{c}{G["ebit"]}', f'={c}{G["ebitda"]}-{c}{G["afa"]}', BLACK, EUR)
        put(gv, f'{c}{G["zins"]}', f'=SUMIF(Fz_Jahr,{c}$4,Fz_Zins)', GREEN, EUR)
        put(gv, f'{c}{G["ebt"]}', f'={c}{G["ebit"]}-{c}{G["zins"]}', BLACK, EUR)
        put(gv, f'{c}{G["vv"]}', 0 if j == 1 else f'={prev}{G["vv"]}-MIN({prev}{G["vv"]},MAX({prev}{G["ebt"]},0))+MAX(0,-{prev}{G["ebt"]})', BLACK, EUR)
        put(gv, f'{c}{G["basis"]}', f'=MAX(0,{c}{G["ebt"]}-{c}{G["vv"]})', BLACK, EUR)
        put(gv, f'{c}{G["steuer"]}', f'={c}{G["basis"]}*Steuersatz', BLACK, EUR)
        put(gv, f'{c}{G["ergebnis"]}', f'={c}{G["ebt"]}-{c}{G["steuer"]}', BOLD, EUR)
        put(gv, f'{c}{G["tilg"]}', f'=SUMIF(Fz_Jahr,{c}$4,Fz_Tilg)', GREEN, EUR)
        put(gv, f'{c}{G["kd"]}', f'={c}{G["zins"]}+{c}{G["tilg"]}', BLACK, EUR)
        put(gv, f'{c}{G["dscr"]}', f'=IF({c}{G["kd"]}>0,({c}{G["ebitda"]}-{c}{G["steuer"]})/{c}{G["kd"]},0)', BOLD, '0.00')
        put(gv, f'{c}{G["cf"]}', f'={c}{G["ebitda"]}-{c}{G["steuer"]}-{c}{G["kd"]}', BLACK, EUR)
        put(gv, f'{c}{G["mitgl"]}', f'=INDEX(Mg_Ende,{c}$4*12)', GREEN, NUM)
        put(gv, f'{c}{G["marge"]}', f'=IF({c}{G["umsatz"]}>0,{c}{G["ebitda"]}/{c}{G["umsatz"]},0)', BLACK, PCT)
        put(gv, f'{c}{G["pq"]}', f'=IF({c}{G["umsatz"]}>0,{c}{G["personal"]}/{c}{G["umsatz"]},0)', BLACK, PCT)
        put(gv, f'{c}{G["rq"]}', f'=IF({c}{G["umsatz"]}>0,{c}{G["raum"]}/{c}{G["umsatz"]},0)', BLACK, PCT)
    names.add_range('GuV_Steuern', gv, f'C{G["steuer"]}:G{G["steuer"]}')
    names.add_range('GuV_EBITDA', gv, f'C{G["ebitda"]}:G{G["ebitda"]}')
    names.add_range('GuV_Ergebnis', gv, f'C{G["ergebnis"]}:G{G["ergebnis"]}')
    names.add_range('GuV_DSCR', gv, f'C{G["dscr"]}:G{G["dscr"]}')
    gv.freeze_panes = 'C5'

    # ======================= Kennzahlen =======================
    kz = wb.create_sheet('Kennzahlen')
    setw(kz, {'A': 3, 'B': 70, 'C': 18, 'D': 50})
    put(kz, 'B1', 'Kennzahlen', H1)
    k_rows = [
        ('Kapitalbedarf gesamt', '=Kapitalbedarf', EUR, ''),
        ('davon Investition (ohne Reserve)', '=Inv_Summe', EUR, ''),
        ('davon Anlaufreserve (angesetzt)', '=Reserve', EUR, 'Eingabe Annahmen'),
        ('Berechneter Reservebedarf 24 Monate (max. kumulierter Fehlbetrag Vorlauf–Monat 24 × (1 + Puffer))', f'=MAX(0,-MIN(OFFSET(Lq_BetriebKum,0,0,1,Vorlauf+24)))*(1+Puffer_Pct)', EUR, 'sollte ≤ angesetzter Reserve sein'),
        ('Reserve ausreichend?', '=IF(Reserve>=C7,"ja","nein – Reserve erhöhen")', '@', ''),
        ('Berechneter Reservebedarf bis Monat 36 (× (1 + Puffer))', f'=MAX(0,-MIN(OFFSET(Lq_BetriebKum,0,0,1,Vorlauf+36)))*(1+Puffer_Pct)', EUR, ''),
        ('Niedrigster Kassenbestand (Liquiditätsplan)', '=MIN(Lq_Ende)', EUR, ''),
        ('Monat des niedrigsten Kassenbestands', '=INDEX(Lq_Monat,MATCH(MIN(Lq_Ende),Lq_Ende,0))', '0', ''),
        ('Eigenkapitalquote (inkl. Nachrang) am Kapitalbedarf', '=EK_Quote', PCT, ''),
        ('Über-/Unterdeckung Mittelherkunft', '=Deckung', EUR, 'soll ≥ 0 sein'),
        ('Ø Beitrag je Mitglied brutto (Mix)', '=ARPU_brutto', EUR2, ''),
        ('Deckungsbeitrag je Mitglied und Monat netto (Beitrag + Service + Nebenumsatz − Wareneinsatz − Zahlungsverkehr)', '=(ARPU_brutto+Service_Halbjahr/6+Neben_Getraenke+Neben_Wellness+Neben_PT+Neben_Merch)/(1+UST)-Neben_Getraenke/(1+UST)*Neben_WE-Payment_Pct*(ARPU_brutto+Service_Halbjahr/6)', EUR2, ''),
        ('Fixkosten je Monat im Monat 36 (ohne Wareneinsatz/Zahlungsverkehr)', f'=INDEX(Ko_Gesamt,MATCH(36,Ko_Monat,0))-INDEX(Ko_WE,MATCH(36,Ko_Monat,0))-INDEX(Kosten!{first_col}{row_of["payment"]}:{last_col}{row_of["payment"]},MATCH(36,Ko_Monat,0))', EUR, ''),
        ('Kapitaldienst je Monat im Monat 36', '=INDEX(Fz_Zins,MATCH(36,Fz_Monat,0))+INDEX(Fz_Tilg,MATCH(36,Fz_Monat,0))', EUR, ''),
        ('Break-even Mitglieder (Betriebskosten, ohne Kapitaldienst)', '=ROUNDUP(C16/C15,0)', NUM, ''),
        ('Break-even Mitglieder (Vollkosten inkl. Kapitaldienst, Monat 36)', '=ROUNDUP((C16+C17)/C15,0)', NUM, ''),
        ('Erster Monat mit positivem Betriebsergebnis (EBITDA)', '=IFERROR(INDEX(Um_Monat,MATCH(1,Hilf_EBITDA,0)),"nicht innerhalb 60 Monaten")', '0', ''),
        ('Erster Monat mit positivem Cashflow nach Kapitaldienst', '=IFERROR(INDEX(Um_Monat,MATCH(1,Hilf_CF,0)),"nicht innerhalb 60 Monaten")', '0', ''),
        ('Mitglieder Ende Jahr 1 / 2 / 3', '=TEXT(INDEX(Mg_Ende,12),"#,##0")&" / "&TEXT(INDEX(Mg_Ende,24),"#,##0")&" / "&TEXT(INDEX(Mg_Ende,36),"#,##0")', '@', ''),
        ('Umsatz Jahr 3 je m² Bruttofläche', '=Umsatz!E27', EUR, ''),
        ('Mitglieder je m² Bruttofläche (Ende Jahr 3)', '=INDEX(Mg_Ende,36)/Flaeche', '0.00', 'Branchenüblich 1,0–1,5'),
        ('Personalkostenquote Jahr 3', f'=GuV!E{G["pq"]}', PCT, ''),
        ('Raumkostenquote Jahr 3', f'=GuV!E{G["rq"]}', PCT, ''),
        ('EBITDA-Marge Jahr 3', f'=GuV!E{G["marge"]}', PCT, ''),
        ('DSCR Jahr 3 / 4 / 5', f'=TEXT(GuV!E{G["dscr"]},"0.00")&" / "&TEXT(GuV!F{G["dscr"]},"0.00")&" / "&TEXT(GuV!G{G["dscr"]},"0.00")', '@', 'Banken erwarten i. d. R. ≥ 1,2 nach Anlaufphase'),
        ('Restschuld Ende Jahr 5', '=INDEX(Fz_Rest,MATCH(60,Fz_Monat,0))', EUR, ''),
        ('Vorsteuer auf Investitionen (Erstattung, Zwischenfinanzierung über Kontokorrent)', '=Inv_Vorsteuer_Gesamt', EUR, ''),
    ]
    put(kz, 'B3', 'Kennzahl', BOLD, fill=HEAD); put(kz, 'C3', 'Wert', BOLD, fill=HEAD); put(kz, 'D3', 'Hinweis', BOLD, fill=HEAD)
    for i, (lab_, f, fmt, note) in enumerate(k_rows):
        put(kz, f'B{4 + i}', lab_); put(kz, f'C{4 + i}', f, BLACK, fmt); put(kz, f'D{4 + i}', note, GREY)
    # Hilfszeilen auf Umsatzblatt: EBITDA>=0 und CF>=0 je Monat
    put(um, 'B30', 'Hilfszeilen für Kennzahlen', H2)
    put(um, 'B31', 'Betriebsergebnis (Umsatz − Kosten) ≥ 0 → 1'); put(um, 'B32', 'Cashflow nach Kapitaldienst ≥ 0 → 1')
    for j, m in enumerate(months_op):
        c = L(3 + j)
        kc = col_of(m)
        put(um, f'{c}31', f'=IF({c}16-Kosten!{kc}{r_sum}>=0,1,0)', BLACK, '0')
        put(um, f'{c}32', f'=IF({c}16-Kosten!{kc}{r_sum}-Finanzierung!{kc}{rr + 1}-Finanzierung!{kc}{rr + 2}>=0,1,0)', BLACK, '0')
    names.add_range('Hilf_EBITDA', um, f'C31:{L(2 + MONATE)}31')
    names.add_range('Hilf_CF', um, f'C32:{L(2 + MONATE)}32')

    # ======================= Szenarien (Werte) =======================
    sz = wb.create_sheet('Szenarien')
    setw(sz, {'A': 3, 'B': 56, 'C': 18, 'D': 18, 'E': 18})
    put(sz, 'B1', 'Szenarienvergleich', H1)
    put(sz, 'B2', 'Werte (keine Formeln): erzeugt durch Umschalten der Zelle „Szenario“ auf dem Blatt Annahmen und Neuberechnung; Stand ' + stand + '. Zum Aktualisieren Szenario 1/2/3 wählen und die Blätter GuV/Liquiditaet/Kennzahlen ablesen.', GREY)
    put(sz, 'B4', 'Kennzahl', BOLD, fill=HEAD); put(sz, 'C4', 'pessimistisch', BOLD, fill=HEAD); put(sz, 'D4', 'Basis', BOLD, fill=HEAD); put(sz, 'E4', 'optimistisch', BOLD, fill=HEAD)
    if szenario_werte:
        keys = [('Neuzugänge-Faktor / Kündigungsquote / Beitragsniveau', 'faktoren', '@'), ('Mitglieder Ende Jahr 1', 'mitglieder_m12', NUM), ('Mitglieder Ende Jahr 2', 'mitglieder_m24', NUM), ('Mitglieder Ende Jahr 3', 'mitglieder_m36', NUM),
                ('Umsatz Jahr 1', 'umsatz_j1', EUR), ('Umsatz Jahr 2', 'umsatz_j2', EUR), ('Umsatz Jahr 3', 'umsatz_j3', EUR), ('EBITDA Jahr 1', 'ebitda_j1', EUR), ('EBITDA Jahr 2', 'ebitda_j2', EUR), ('EBITDA Jahr 3', 'ebitda_j3', EUR),
                ('Jahresergebnis Jahr 1', 'ergebnis_j1', EUR), ('Jahresergebnis Jahr 2', 'ergebnis_j2', EUR), ('Jahresergebnis Jahr 3', 'ergebnis_j3', EUR), ('Jahresergebnis Jahr 5', 'ergebnis_j5', EUR),
                ('DSCR Jahr 3', 'dscr_j3', '0.00'), ('DSCR Jahr 4', 'dscr_j4', '0.00'), ('Niedrigster Kassenbestand (Monat)', 'min_kasse_txt', '@'), ('Reservebedarf 24 Monate inkl. Puffer', 'reserve_bedarf_24', EUR),
                ('Erster Monat EBITDA ≥ 0', 'be_monat_ebitda', '0'), ('Erster Monat Cashflow nach Kapitaldienst ≥ 0', 'be_monat_cf', '0'), ('Break-even Mitglieder (Vollkosten M36)', 'be_mitglieder', NUM)]
        for i, (lab_, key, fmt) in enumerate(keys):
            put(sz, f'B{5 + i}', lab_)
            for col, s in zip('CDE', ('pessimistisch', 'basis', 'optimistisch')):
                v = szenario_werte[s].get(key)
                put(sz, f'{col}{5 + i}', v if v is not None else 'n. e.', BLACK, fmt)

    # Druckeinstellungen
    for w in wb.worksheets:
        w.page_setup.orientation = 'landscape'
        w.page_setup.fitToWidth = 1
        w.page_setup.fitToHeight = 0
        w.sheet_properties.pageSetUpPr.fitToPage = True
        w.print_options.gridLines = False
    wb.save(path)


if __name__ == '__main__':
    import sys
    a = M.load(sys.argv[1]); d = M.load(sys.argv[2])
    write_excel(a, d, sys.argv[3])
    print('geschrieben', sys.argv[3])
