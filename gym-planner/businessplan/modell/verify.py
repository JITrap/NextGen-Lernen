"""Vergleicht die per LibreOffice neu berechnete Excel-Datei mit dem Python-Spiegel (model.compute)."""
from __future__ import annotations

import json
import subprocess
import sys

from openpyxl import load_workbook

import model as M

RECALC = '/root/.claude/skills/synced/8de5c00b-c1fc-4df8-9ee2-8554408044e0_9e49e0d6-3e87-4558-a77c-8158fbd2fca1/xlsx/scripts/recalc.py'


def recalc(path: str) -> dict:
    out = subprocess.run([sys.executable, RECALC, path, '120'], capture_output=True, text=True)
    try:
        return json.loads(out.stdout.strip().splitlines()[-1])
    except Exception:
        return {'error': out.stdout + out.stderr}


def cmp(name: str, xl: float | None, py: float, tol: float = 1.0, rel: float = 0.0005) -> str | None:
    if xl is None:
        return f'{name}: Excel leer, Python {py:.2f}'
    if abs(xl - py) > max(tol, abs(py) * rel):
        return f'{name}: Excel {xl:.2f} ≠ Python {py:.2f} (Δ {xl - py:.2f})'
    return None


def verify(a: dict, daten: dict, path: str, szenario: str = 'basis') -> list[str]:
    r = M.compute(a, daten, szenario)
    wb = load_workbook(path, data_only=True)
    probs: list[str] = []
    V = a['projekt']['vorlauf_monate']
    monate = r['monate']

    def col(ws, m: int) -> int:
        # Spalte C = erster Monat der jeweiligen Zeitachse
        row4 = [c.value for c in ws[4]]
        for i, v in enumerate(row4):
            if v == m and i >= 2:
                return i + 1
        raise KeyError(f'Monat {m} nicht in {ws.title}')

    iv = wb['Investition']
    # Kapitalbedarf/Investition
    for row in iv.iter_rows(min_row=1, max_row=60, min_col=3, max_col=4):
        lab, val = row[0].value, row[1].value
        if lab == 'Investitionssumme (ohne Anlaufreserve)':
            probs.append(cmp('Investitionssumme', val, r['inv_summe']))
        if lab == 'Kapitalbedarf gesamt':
            probs.append(cmp('Kapitalbedarf', val, r['kapitalbedarf']))
        if lab and lab.startswith('Gerätewert'):
            probs.append(cmp('Gerätewert', val, r['geraetewert']))
    mg = wb['Mitglieder']
    um = wb['Umsatz']
    ko = wb['Kosten']
    fz = wb['Finanzierung']
    lq = wb['Liquiditaet']
    gv = wb['GuV']
    for m in (1, 2, 6, 12, 13, 24, 25, 36, 48, 60):
        c = col(mg, m)
        probs.append(cmp(f'Mitglieder Ende M{m}', mg.cell(10, c).value, r['bestand'][m], 0.5))
        probs.append(cmp(f'Mitglieder Anfang M{m}', mg.cell(7, c).value, r['anfang'][m], 0.5))
        c = col(um, m)
        probs.append(cmp(f'Umsatz M{m}', um.cell(16, c).value, r['umsatz'][m]))
        probs.append(cmp(f'Beitrag M{m}', um.cell(9, c).value, r['beitrag'][m]))
        probs.append(cmp(f'Service M{m}', um.cell(10, c).value, r['service'][m]))
        probs.append(cmp(f'Aufnahme M{m}', um.cell(11, c).value, r['aufnahme'][m]))
    # Kosten: Summenzeile suchen
    sum_row = next(i for i in range(1, 60) if ko.cell(i, 2).value == 'Kosten gesamt (netto)')
    pers_row = next(i for i in range(1, 60) if ko.cell(i, 2).value and ko.cell(i, 2).value.startswith('Personal'))
    for m in monate:
        if m in (-V, -3, -1, 1, 3, 6, 12, 13, 24, 25, 36, 60):
            c = col(ko, m)
            probs.append(cmp(f'Kosten M{m}', ko.cell(sum_row, c).value, r['kosten_summe'][m]))
            probs.append(cmp(f'Personal M{m}', ko.cell(pers_row, c).value, r['kosten'][m]['personal']))
            probs.append(cmp(f'Vorsteuer Kosten M{m}', ko.cell(sum_row + 2, c).value, r['vst_kosten'][m]))
    # Finanzierung: Summenzeilen
    z_row = next(i for i in range(1, 80) if fz.cell(i, 2).value == 'Summe Zinsen')
    for m in monate:
        if m in (-4, -1, 1, 12, 24, 25, 36, 60):
            c = col_fz = None
            row4 = None
            # Monatszeile der Tilgungspläne finden
            for i in range(1, 80):
                if fz.cell(i, 2).value == 'Monat' and fz.cell(i, 3).value == -V:
                    row4 = i
                    break
            for j in range(3, 3 + len(monate)):
                if fz.cell(row4, j).value == m:
                    c = j
                    break
            probs.append(cmp(f'Zinsen M{m}', fz.cell(z_row, c).value, r['zins'][m]))
            probs.append(cmp(f'Tilgung M{m}', fz.cell(z_row + 1, c).value, r['tilg'][m]))
            probs.append(cmp(f'Restschuld M{m}', fz.cell(z_row + 3, c).value, r['saldo_ende'][m]))
    # Liquidität
    ende_row = next(i for i in range(1, 40) if lq.cell(i, 2).value == 'Kassenbestand Ende')
    ust_row = next(i for i in range(1, 40) if lq.cell(i, 2).value == 'USt-Zahllast an das Finanzamt')
    erst_row = next(i for i in range(1, 40) if lq.cell(i, 2).value == 'USt-Erstattung vom Finanzamt')
    kum_row = next(i for i in range(1, 40) if lq.cell(i, 2).value == 'kumuliert (Basis Reservebedarf)')
    for m in monate:
        if m in (-V, -4, -3, -2, -1, 1, 2, 3, 6, 12, 18, 24, 30, 36, 48, 60):
            c = col(lq, m)
            probs.append(cmp(f'Kasse Ende M{m}', lq.cell(ende_row, c).value, r['liq'][m]['ende']))
            probs.append(cmp(f'USt-Zahlung M{m}', (lq.cell(ust_row, c).value or 0) - (lq.cell(erst_row, c).value or 0), r['ust_zahlung'][m]))
            probs.append(cmp(f'Betrieb kum M{m}', lq.cell(kum_row, c).value, r['liq'][m]['betrieb_kum']))
    # GuV
    labels = {gv.cell(i, 2).value: i for i in range(1, 40) if gv.cell(i, 2).value}
    keymap = {'Umsatzerlöse (netto)': 'umsatz', 'Personalkosten': 'personal', 'Raumkosten (Miete, Betriebskosten, Energie, Instandhaltung)': 'raumkosten', 'Marketing': 'marketing', 'Sonstige betriebliche Aufwendungen': 'sonstige', 'EBITDA (Betriebsergebnis vor Abschreibung und Zinsen)': 'ebitda', 'Abschreibungen': 'afa', 'Zinsaufwand': 'zinsen', 'Ergebnis vor Steuern (EBT)': 'ebt', 'Ertragsteuern (GewSt, KSt, Soli)': 'steuern', 'Jahresergebnis': 'jahresergebnis', 'Tilgung': 'tilgung', 'Kapitaldienst (Zins + Tilgung)': 'kapitaldienst', 'Verlustvortrag Anfang': 'verlustvortrag_anfang', 'Steuerliche Bemessungsgrundlage': 'steuerbasis', 'Wareneinsatz': 'wareneinsatz'}
    for lab, key in keymap.items():
        for j in range(1, 6):
            probs.append(cmp(f'GuV {key} J{j}', gv.cell(labels[lab], 2 + j).value, r['guv'][j][key]))
    kz = wb['Kennzahlen']
    kzl = {kz.cell(i, 2).value: kz.cell(i, 3).value for i in range(1, 40) if kz.cell(i, 2).value}
    probs.append(cmp('Reservebedarf 24', kzl['Berechneter Reservebedarf 24 Monate (max. kumulierter Fehlbetrag Vorlauf–Monat 24 × (1 + Puffer))'], r['reserve_bedarf_24']))
    probs.append(cmp('Break-even Mitglieder (Vollkosten)', kzl['Break-even Mitglieder (Vollkosten inkl. Kapitaldienst, Monat 36)'], r['be_mitglieder'], 1))
    probs.append(cmp('Break-even Mitglieder (ohne KD)', kzl['Break-even Mitglieder (Betriebskosten, ohne Kapitaldienst)'], r['be_mitglieder_ohne_kd'], 1))
    probs.append(cmp('Min Kasse', kzl['Niedrigster Kassenbestand (Liquiditätsplan)'], r['min_kasse'][1]))
    be = kzl['Erster Monat mit positivem Betriebsergebnis (EBITDA)']
    if isinstance(be, (int, float)) or r['be_monat_ebitda'] is not None:
        if be != r['be_monat_ebitda']:
            probs.append(f'Break-even Monat EBITDA: Excel {be} ≠ Python {r["be_monat_ebitda"]}')
    return [p for p in probs if p]


if __name__ == '__main__':
    a = M.load(sys.argv[1]); d = M.load(sys.argv[2]); path = sys.argv[3]
    res = recalc(path)
    print('recalc:', json.dumps({k: res.get(k) for k in ('status', 'total_formulas', 'total_errors')}), (res.get('error_summary') or res.get('error') or '')[:2000] if res.get('total_errors') or res.get('error') else '')
    probs = verify(a, d, path, sys.argv[4] if len(sys.argv) > 4 else 'basis')
    print('Abweichungen:', len(probs))
    for p in probs[:60]:
        print(' -', p)
